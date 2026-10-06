from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.reports.export.xlsx.analytics_builder import (
    build_analytics_workbook,
    validate_analytics_xlsx_artifact,
)
from app.reports.export.xlsx.analytics_profiles import (
    ANALYTICS_WORKBOOK_PROFILES,
    ANALYTICS_XLSX_REPORT_TYPES,
    prepare_analytics_workbook_data,
)
from app.reports.export.xlsx.analytics_schema import (
    ANALYTICS_WORKBOOK_STANDARD,
    AnalyticsExportContext,
)
from app.reports.export.xlsx.schema import XLSX_REPORTING_STANDARD
from app.reports.export.xlsx.style import KAORI_SLATE_FILL

ANALYTICS_TYPES = {
    'analytics_overview',
    'resource_utilization',
    'capacity_analysis',
    'storage_analysis',
    'workload_analysis',
    'resource_forecast',
    'top_consumers',
    'alerts_anomalies',
    'object_storage_analysis',
    'tenant_governance',
}


def _context(report_type: str) -> AnalyticsExportContext:
    return AnalyticsExportContext(
        report_id='KAORI-TEST-ANX',
        report_name='Analytics XLSX Test',
        report_type=report_type,
        generated_at=datetime(2026, 8, 26, 0, 30, tzinfo=UTC),
        generated_by='Platform Operations',
        scope_label='Kubernetes Cluster',
        scope_value='pgn-k8s-staging',
        namespace_scope='All Namespaces',
        period_label='24h',
        data_cutoff=datetime(2026, 8, 26, 0, 0, tzinfo=UTC),
        classification='Internal',
    )


def test_analytics_xlsx_profiles_cover_every_published_analytics_category() -> None:
    assert ANALYTICS_WORKBOOK_STANDARD == 'kaori-analytics-xlsx-v1.0'
    assert XLSX_REPORTING_STANDARD == 'kaori-reporting-xlsx-v2.0'
    assert ANALYTICS_XLSX_REPORT_TYPES == ANALYTICS_TYPES
    assert set(ANALYTICS_WORKBOOK_PROFILES) == ANALYTICS_TYPES

    for report_type, profile in ANALYTICS_WORKBOOK_PROFILES.items():
        assert profile.report_type == report_type
        assert profile.title
        assert profile.description
        assert len(profile.metric_specs) == 4
        assert profile.sheet_names
        assert profile.primary_visual in {
            'line',
            'ranked_bar',
            'composition_bar',
            'capacity_bar',
            'comparison_bar',
        }


def test_resource_utilization_preparation_preserves_committed_metrics_trends_and_consumers() -> None:
    payload = {
        'summary': {
            'cpu_utilization': 18.2,
            'memory_utilization': 61.4,
            'storage_utilization': 42.5,
            'pod_utilization': 53.0,
        },
        'cluster_info': {'total_nodes': 5, 'ready_nodes': 5},
        'trends': {
            'cpu': [{'value': 14.0}, {'value': 18.2}],
            'memory': [{'value': 58.1}, {'value': 61.4}],
            'pods': [{'value': 49.0}, {'value': 53.0}],
        },
        'top_cpu_consumers': [
            {'namespace': 'monitoring', 'pod': 'prometheus-0', 'value': 0.71},
        ],
        'top_memory_consumers': [
            {'namespace': 'openebs', 'pod': 'cstor-0', 'value': 16.4},
        ],
        'source_evidence': [{'source': 'Committed Analytics', 'freshness': 'Current'}],
    }

    data = prepare_analytics_workbook_data('resource_utilization', payload, _context('resource_utilization'))

    assert [metric.label for metric in data.metrics] == [
        'CPU Utilization',
        'Memory Utilization',
        'Storage Utilization',
        'Pod Utilization',
    ]
    assert [metric.value for metric in data.metrics] == [18.2, 61.4, 42.5, 53.0]
    assert data.primary_chart is not None
    assert data.primary_chart.visual_type == 'line'
    assert data.primary_chart.source_sheet == 'Trends'
    assert [series.label for series in data.primary_chart.series] == ['CPU', 'Memory', 'Pods']
    assert {'Analysis', 'Trends', 'Top Consumers'} <= {table.sheet_name for table in data.tables}
    trends = next(table for table in data.tables if table.sheet_name == 'Trends')
    assert trends.rows[-1]['memory_percent'] == 61.4
    assert '61.4%' in data.readout


def test_resource_utilization_line_chart_uses_timestamps_only_available_series_and_sparse_presentation(
    tmp_path: Path,
) -> None:
    trend_points = []
    for index in range(12):
        trend_points.append(
            {
                'timestamp': f'2026-08-26T{index:02d}:00:00+07:00',
                'cpu': 1.0 + (index * 0.1),
                'memory': 55.0 + (index * 0.2),
                'pods': 46.0 + (index * 0.1),
            }
        )
    payload = {
        'summary': {
            'cpu_utilization': 2.1,
            'memory_utilization': 57.2,
            'storage_utilization': None,
            'pod_utilization': 47.1,
        },
        'trends': {
            'cpu': [{'timestamp': row['timestamp'], 'value': row['cpu']} for row in trend_points],
            'memory': [{'timestamp': row['timestamp'], 'value': row['memory']} for row in trend_points],
            'pods': [{'timestamp': row['timestamp'], 'value': row['pods']} for row in trend_points],
            'storage': [],
        },
    }

    data = prepare_analytics_workbook_data('resource_utilization', payload, _context('resource_utilization'))
    assert data.primary_chart is not None
    assert data.primary_chart.category_key == 'timestamp'
    assert [series.label for series in data.primary_chart.series] == ['CPU', 'Memory', 'Pods']

    output_path = tmp_path / 'resource-utilization-sparse.xlsx'
    build_analytics_workbook(
        report_type='resource_utilization',
        payload=payload,
        context=_context('resource_utilization'),
        output_path=output_path,
    )
    validate_analytics_xlsx_artifact(output_path)

    workbook = load_workbook(output_path)
    overview = workbook['Overview']
    chart = overview._charts[0]
    assert chart.__class__.__name__ == 'LineChart'
    assert [series.tx.v for series in chart.series] == ['CPU', 'Memory', 'Pods']
    assert all(series.marker.symbol is None for series in chart.series)
    with ZipFile(output_path) as archive:
        chart_xml = archive.read('xl/charts/chart1.xml').decode('utf-8')
    chart_root = ElementTree.fromstring(chart_xml)
    chart_namespace = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    marker_symbols = chart_root.findall('.//c:marker/c:symbol', chart_namespace)
    assert [symbol.get('val') for symbol in marker_symbols].count('none') == 3
    assert chart.y_axis.scaling.min == 0
    assert chart.y_axis.scaling.max == 100
    assert chart.y_axis.majorUnit == 25
    assert chart.y_axis.majorGridlines is None
    assert chart.y_axis.title is None
    assert chart.x_axis.title is None
    assert overview['J15'].value == 'Observed At'
    assert overview['K15'].value == 'CPU'
    assert overview['L15'].value == 'Memory'
    assert overview['M15'].value == 'Pods'
    assert overview['N15'].value is None
    assert overview.column_dimensions['J'].width >= 22
    assert overview.column_dimensions['K'].width >= 14


def test_analytics_overview_uses_current_vs_projected_resource_posture() -> None:
    payload = {
        'resource_snapshot': {
            'cpu': {'percent': 22.4, 'status': 'Normal'},
            'memory': {'percent': 61.4, 'status': 'Warning'},
            'pods': {'percent': 51.0, 'status': 'Normal'},
            'storage': {'percent': 44.0, 'status': 'Normal'},
        },
        'forecast_snapshot': {
            'highest_projected_resource': 'Memory',
            'highest_projected_percent': 67.7,
            'series': [
                {
                    'key': 'cpu',
                    'label': 'CPU',
                    'latest_percent': 22.4,
                    'projected_percent': 31.8,
                    'trend': 'Increasing',
                    'confidence': 'High',
                },
                {
                    'key': 'memory',
                    'label': 'Memory',
                    'latest_percent': 61.4,
                    'projected_percent': 67.7,
                    'trend': 'Increasing',
                    'confidence': 'High',
                },
            ],
        },
        'alerts_snapshot': {'critical': 0, 'warning': 1, 'info': 0},
    }

    data = prepare_analytics_workbook_data('analytics_overview', payload, _context('analytics_overview'))

    assert data.primary_chart is not None
    assert data.primary_chart.visual_type == 'comparison_bar'
    assert data.primary_chart.title == 'Current vs Projected Resource Posture'
    assert [series.label for series in data.primary_chart.series] == ['Current', 'Projected']
    assert data.primary_chart.rows == (
        {
            'resource': 'CPU',
            'current_percent': 22.4,
            'projected_percent': 31.8,
            'change_pp': 9.4,
            'trend': 'Increasing',
            'confidence': 'High',
        },
        {
            'resource': 'Memory',
            'current_percent': 61.4,
            'projected_percent': 67.7,
            'change_pp': 6.3,
            'trend': 'Increasing',
            'confidence': 'High',
        },
    )
    trends = next(table for table in data.tables if table.sheet_name == 'Trends')
    assert [column.key for column in trends.columns] == [
        'resource',
        'current_percent',
        'projected_percent',
        'change_pp',
        'trend',
        'confidence',
    ]
    assert trends.rows[1]['projected_percent'] == 67.7
    assert '67.7%' in data.readout


def test_analytics_overview_does_not_create_visual_from_empty_forecast_values() -> None:
    payload = {
        'resource_snapshot': {
            'memory': {'percent': 61.4},
        },
        'forecast_snapshot': {
            'highest_projected_percent': None,
            'series': [
                {
                    'key': 'memory',
                    'label': 'Memory',
                    'latest_percent': None,
                    'projected_percent': None,
                    'trend': 'Unknown',
                    'confidence': 'Not Available',
                }
            ],
        },
    }

    data = prepare_analytics_workbook_data('analytics_overview', payload, _context('analytics_overview'))

    assert data.primary_chart is None
    assert all(table.sheet_name != 'Trends' for table in data.tables)


def test_alerts_preparation_uses_composition_severity_and_accounts_for_unclassified_alerts() -> None:
    payload = {
        'summary': {
            'active_alerts': 10,
            'critical_alerts': 1,
            'warning_alerts': 5,
            'info_alerts': 2,
            'detected_anomalies': 3,
        },
        'alerts': [
            {'name': 'HighMemory', 'severity': 'critical', 'status': 'firing', 'namespace': 'payments'},
            {'name': 'RestartBurst', 'severity': 'warning', 'status': 'firing', 'namespace': 'core'},
        ],
        'anomalies': [{'name': 'Pod Restart Spike', 'severity': 'warning', 'namespace': 'core'}],
        'actionable_recommendations': [],
    }

    data = prepare_analytics_workbook_data('alerts_anomalies', payload, _context('alerts_anomalies'))

    assert data.primary_chart is not None
    assert data.primary_chart.visual_type == 'composition_bar'
    assert data.primary_chart.rows == (
        {'category': 'Critical', 'value': 1},
        {'category': 'Warning', 'value': 5},
        {'category': 'Info', 'value': 2},
        {'category': 'Unclassified', 'value': 2},
    )
    assert 'Recommendations' not in {table.sheet_name for table in data.tables}


def test_object_storage_preparation_owns_its_category_and_ranked_consumer_visual() -> None:
    payload = {
        'system': {'name': 'minio-pgn-development', 'environment': 'Development'},
        'summary': {
            'capacity': '120 GiB',
            'usage': '76 GiB',
            'headroom': '44 GiB',
            'utilization': '63.3%',
            'buckets': '5',
            'objects': '12,300',
        },
        'capacity_usage': {
            'capacity_total': '120 GiB',
            'usage_total': '76 GiB',
            'capacity_free': '44 GiB',
            'utilization': '63.3%',
        },
        'top_consumers': [
            {'name': 'thanos', 'usage_gib': 64.25, 'object_count': 9000},
            {'name': 'crm', 'usage_gib': 7.37, 'object_count': 2100},
        ],
        'performance': {'request_rate': '125 req/s', 'error_ratio': '0.2%', 'traffic': '32 MiB/s'},
        'alerts': [],
    }

    context = _context('object_storage_analysis')
    context = replace(context, scope_label='Object Storage System', scope_value='minio-pgn-development')
    data = prepare_analytics_workbook_data('object_storage_analysis', payload, context)

    assert data.primary_chart is not None
    assert data.primary_chart.visual_type == 'ranked_bar'
    assert data.primary_chart.source_sheet == 'Top Consumers'
    assert data.primary_chart.rows[0]['category'] == 'thanos'
    assert data.primary_chart.rows[0]['value'] == 64.25
    assert {'Analysis', 'Top Consumers'} <= {table.sheet_name for table in data.tables}


def test_resource_utilization_workbook_is_professional_and_runtime_valid(tmp_path: Path) -> None:
    payload = {
        'summary': {
            'cpu_utilization': 18.2,
            'memory_utilization': 61.4,
            'storage_utilization': 42.5,
            'pod_utilization': 53.0,
        },
        'trends': {
            'cpu': [{'value': 14.0}, {'value': 16.0}, {'value': 18.2}],
            'memory': [{'value': 58.1}, {'value': 60.0}, {'value': 61.4}],
            'pods': [{'value': 49.0}, {'value': 51.0}, {'value': 53.0}],
        },
        'top_cpu_consumers': [{'namespace': 'monitoring', 'pod': 'prometheus-0', 'value': 0.71}],
        'top_memory_consumers': [{'namespace': 'openebs', 'pod': 'cstor-0', 'value': 16.4}],
        'source_evidence': [{'source': 'Committed Analytics', 'freshness': 'Current'}],
    }
    output_path = tmp_path / 'resource-utilization.xlsx'

    build_analytics_workbook(
        report_type='resource_utilization',
        payload=payload,
        context=_context('resource_utilization'),
        output_path=output_path,
    )
    validate_analytics_xlsx_artifact(output_path)

    workbook = load_workbook(output_path)
    assert workbook.sheetnames[0] == 'Overview'
    assert workbook.sheetnames[-2:] == ['Metadata', 'Field Dictionary']
    assert {'Analysis', 'Trends', 'Top Consumers'} <= set(workbook.sheetnames)
    overview = workbook['Overview']
    assert overview['A1'].value == 'KAORI — Resource Utilization Analytics'
    assert overview['A9'].value == 'Executive Readout'
    assert len(overview._charts) == 1
    assert overview._charts[0].__class__.__name__ == 'LineChart'
    assert [series.tx.v for series in overview._charts[0].series] == ['CPU', 'Memory', 'Pods']
    assert overview['A32'].hyperlink.target == "#'Analysis'!A1"
    assert overview['A33'].hyperlink is not None
    assert workbook['Analysis']['A3'].hyperlink.target == "#'Overview'!A1"
    metadata_values = {
        workbook['Metadata'].cell(row=row, column=1).value: workbook['Metadata'].cell(row=row, column=2).value
        for row in range(1, workbook['Metadata'].max_row + 1)
    }
    assert metadata_values['Workbook Standard'] == ANALYTICS_WORKBOOK_STANDARD
    assert metadata_values['XLSX Standard'] == XLSX_REPORTING_STANDARD


def test_capacity_workbook_uses_resource_categories_and_stacked_utilization_headroom(tmp_path: Path) -> None:
    payload = {
        'capacity_summary': {
            'cpu': {'usage_percent': 31.2},
            'memory': {'usage_percent': 67.5},
            'pods': {'usage_percent': 49.0},
            'storage': {'usage_percent': 54.3},
        },
    }
    output_path = tmp_path / 'capacity-analysis.xlsx'
    build_analytics_workbook(
        report_type='capacity_analysis',
        payload=payload,
        context=_context('capacity_analysis'),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    overview = workbook['Overview']
    chart = overview._charts[0]
    assert chart.__class__.__name__ == 'BarChart'
    assert chart.grouping == 'stacked'
    assert chart.overlap == 100
    assert chart.legend is not None
    assert chart.legend.position == 'b'
    assert [series.tx.v for series in chart.series] == ['Utilization', 'Headroom']
    assert chart.dLbls is not None
    assert chart.dLbls.showVal is False
    assert chart.dLbls.showLegendKey is False
    assert chart.dLbls.showCatName is False
    assert chart.dLbls.showSerName is False
    assert chart.dLbls.showPercent is False
    assert chart.x_axis.delete is False
    assert chart.x_axis.tickLblPos == 'nextTo'
    assert chart.x_axis.title is None
    assert chart.y_axis.scaling.min == 0
    assert chart.y_axis.scaling.max == 100
    assert chart.y_axis.majorUnit == 25
    assert chart.y_axis.majorGridlines is None
    assert chart.y_axis.title is None
    assert overview.column_dimensions['J'].width >= 20
    assert overview.column_dimensions['K'].width >= 16
    assert overview.column_dimensions['L'].width >= 16
    assert overview['J15'].value == 'Resource'
    assert overview['K15'].value == 'Utilization'
    assert overview['L15'].value == 'Headroom'
    assert overview['J16'].value == 'CPU'
    assert overview['K16'].value == 31.2
    assert overview['L16'].value == 68.8


def test_object_storage_workbook_uses_category_title_and_professional_ranked_support(tmp_path: Path) -> None:
    payload = {
        'summary': {'capacity': '120 GiB', 'usage': '76 GiB', 'headroom': '44 GiB', 'utilization': '63.3%'},
        'capacity_usage': {
            'capacity_total': '120 GiB',
            'usage_total': '76 GiB',
            'capacity_free': '44 GiB',
            'utilization': '63.3%',
        },
        'top_consumers': [
            {'name': 'thanos', 'usage_gib': 64.25, 'object_count': 9000},
            {'name': 'crm', 'usage_gib': 7.37, 'object_count': 2100},
        ],
    }
    output_path = tmp_path / 'object-storage-analysis.xlsx'
    context = replace(
        _context('object_storage_analysis'), scope_label='Object Storage System', scope_value='minio-pgn-development'
    )
    build_analytics_workbook(
        report_type='object_storage_analysis', payload=payload, context=context, output_path=output_path
    )

    workbook = load_workbook(output_path)
    overview = workbook['Overview']
    assert overview['A1'].value == 'KAORI — Object Storage Analytics'
    assert overview['J15'].value == 'Rank'
    assert overview['K15'].value == 'Category'
    assert overview['M15'].value == 'Share of Total'
    assert overview.column_dimensions['K'].width >= 34
    assert overview.column_dimensions['L'].width >= 20
    assert 'Not Available' not in str(overview['L16'].number_format)


def test_ranked_bar_uses_value_only_labels_short_chart_categories_and_full_supporting_names(tmp_path: Path) -> None:
    long_name = 'prometheus-prometheus-kube-prometheus-prometheus-0'
    payload = {
        'consumer_payload': {
            'rows': [
                {
                    'resource': 'Memory',
                    'namespace': 'monitoring',
                    'name': long_name,
                    'current_usage': '17.34 GiB',
                    'current_display': '17.34 GiB',
                },
                {
                    'resource': 'Memory',
                    'namespace': 'openebs',
                    'name': 'cstor-disk-dfgm-7f4dfcd7cd',
                    'current_usage': '16.68 GiB',
                    'current_display': '16.68 GiB',
                },
            ]
        }
    }
    output_path = tmp_path / 'ranked-bar-hardening.xlsx'
    build_analytics_workbook(
        report_type='top_consumers',
        payload=payload,
        context=_context('top_consumers'),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    overview = workbook['Overview']
    chart = overview._charts[0]
    assert chart.__class__.__name__ == 'BarChart'
    assert chart.legend is None
    assert chart.dLbls is not None
    assert chart.dLbls.showVal is True
    assert chart.dLbls.showLegendKey is False
    assert chart.dLbls.showCatName is False
    assert chart.dLbls.showSerName is False
    assert chart.dLbls.showPercent is False
    assert chart.y_axis.title is None
    assert chart.y_axis.majorGridlines is None
    assert chart.visible_cells_only is False
    assert overview['K16'].value == long_name
    assert overview['O16'].value.endswith('…')
    assert len(overview['O16'].value) <= 32
    assert overview.column_dimensions['O'].hidden is True
    assert 'Not Available' not in str(overview['L16'].number_format)


def test_composition_bar_renders_numeric_supporting_evidence_without_blank_primary_visual(tmp_path: Path) -> None:
    payload = {
        'tenant_quota_summary': {
            'scope_total': 30,
            'scope_covered': 14,
            'rows': [
                {'namespace': 'payments', 'status': 'Covered'},
                {'namespace': 'core', 'status': 'Not Covered'},
            ],
        },
    }
    output_path = tmp_path / 'composition-bar.xlsx'
    build_analytics_workbook(
        report_type='workload_analysis',
        payload=payload,
        context=_context('workload_analysis'),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    overview = workbook['Overview']
    assert len(overview._charts) == 1
    chart = overview._charts[0]
    assert chart.__class__.__name__ == 'BarChart'
    assert chart.grouping == 'clustered'
    assert chart.legend is None
    assert chart.dLbls is not None
    assert chart.dLbls.showVal is True
    assert chart.dLbls.showCatName is False
    assert chart.dLbls.numFmt == '0.0"%"'
    assert chart.y_axis.scaling.min == 0
    assert chart.y_axis.scaling.max == 100
    assert overview['J15'].value == 'Rank'
    assert overview['M15'].value == 'Share of Total'
    assert overview['M16'].value == pytest.approx(46.6666666667)
    assert overview['M17'].value == pytest.approx(53.3333333333)


def test_formula_like_analytics_values_are_stored_as_text_not_formulas(tmp_path: Path) -> None:
    payload = {
        'consumer_payload': {
            'rows': [
                {
                    'resource': 'CPU',
                    'namespace': '=HYPERLINK("https://bad")',
                    'name': '+cmd',
                    'value': 1.2,
                    'unit': 'cores',
                },
            ]
        }
    }
    output_path = tmp_path / 'top-consumers.xlsx'
    build_analytics_workbook(
        report_type='top_consumers',
        payload=payload,
        context=_context('top_consumers'),
        output_path=output_path,
    )
    workbook = load_workbook(output_path, data_only=False)
    sheet = workbook['Top Consumers']
    values = [cell.value for row in sheet.iter_rows() for cell in row]
    assert '\'=HYPERLINK("https://bad")' in values
    assert "'+cmd" in values


def test_runtime_validator_rejects_missing_primary_chart_when_chartable_evidence_exists(tmp_path: Path) -> None:
    payload = {
        'summary': {
            'cpu_utilization': 18.2,
            'memory_utilization': 61.4,
            'storage_utilization': 42.5,
            'pod_utilization': 53.0,
        },
        'trends': {
            'cpu': [{'value': 14.0}, {'value': 18.2}],
            'memory': [{'value': 58.1}, {'value': 61.4}],
        },
    }
    output_path = tmp_path / 'missing-chart.xlsx'
    build_analytics_workbook(
        report_type='resource_utilization',
        payload=payload,
        context=_context('resource_utilization'),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    workbook['Overview']._charts.clear()
    workbook.save(output_path)

    with pytest.raises(ValueError, match='Primary Visual'):
        validate_analytics_xlsx_artifact(output_path)


def test_runtime_validator_rejects_comparison_bar_without_executive_axis_contract(tmp_path: Path) -> None:
    output_path = tmp_path / 'comparison-axis-regression.xlsx'
    build_analytics_workbook(
        report_type='analytics_overview',
        payload=_representative_payload('analytics_overview'),
        context=_context('analytics_overview'),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    chart = workbook['Overview']._charts[0]
    chart.x_axis.tickLblPos = 'low'
    chart.y_axis.title = 'Percent'
    chart.y_axis.majorUnit = None
    workbook.save(output_path)

    with pytest.raises(ValueError, match='executive axis presentation'):
        validate_analytics_xlsx_artifact(output_path)


def test_runtime_validator_rejects_line_chart_without_sparse_percent_presentation(tmp_path: Path) -> None:
    output_path = tmp_path / 'line-presentation-regression.xlsx'
    build_analytics_workbook(
        report_type='resource_utilization',
        payload=_representative_payload('resource_utilization'),
        context=_context('resource_utilization'),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    chart = workbook['Overview']._charts[0]
    chart.y_axis.title = 'Utilization (%)'
    chart.y_axis.scaling.max = None
    chart.series[0].marker.symbol = 'circle'
    workbook.save(output_path)

    with pytest.raises(ValueError, match='sparse percent presentation'):
        validate_analytics_xlsx_artifact(output_path)


def _representative_payload(report_type: str) -> dict[str, object]:
    payloads: dict[str, dict[str, object]] = {
        'analytics_overview': {
            'resource_snapshot': {
                'cpu': {'percent': 22.0},
                'memory': {'percent': 64.0},
                'pods': {'percent': 51.0},
                'storage': {'percent': 44.0},
            },
            'forecast_snapshot': {
                'highest_projected_resource': 'Memory',
                'highest_projected_percent': 72.0,
                'series': [
                    {
                        'key': 'cpu',
                        'label': 'CPU',
                        'latest_percent': 22.0,
                        'projected_percent': 29.0,
                        'trend': 'Increasing',
                        'confidence': 'High',
                    },
                    {
                        'key': 'memory',
                        'label': 'Memory',
                        'latest_percent': 64.0,
                        'projected_percent': 72.0,
                        'trend': 'Increasing',
                        'confidence': 'High',
                    },
                ],
            },
            'alerts_snapshot': {'critical': 1, 'warning': 2, 'info': 1},
        },
        'resource_utilization': {
            'summary': {
                'cpu_utilization': 18.2,
                'memory_utilization': 61.4,
                'storage_utilization': 42.5,
                'pod_utilization': 53.0,
            },
            'trends': {
                'cpu': [{'value': 14.0}, {'value': 18.2}],
                'memory': [{'value': 58.1}, {'value': 61.4}],
                'pods': [{'value': 49.0}, {'value': 53.0}],
            },
        },
        'capacity_analysis': {
            'capacity_summary': {
                'cpu': {'usage_percent': 32.0},
                'memory': {'usage_percent': 58.0},
                'pods': {'usage_percent': 47.0},
                'storage': {'usage_percent': 41.0},
            },
        },
        'storage_analysis': {
            'storage_summary': {
                'persistent_storage': {
                    'pv': {'count': 18, 'capacity': '920 GiB'},
                    'pvc': {'count': 21, 'used_percent': 63.0, 'used': '580 GiB', 'free': '340 GiB'},
                },
                'node_filesystem': {'rows': [{'name': '/var/lib/kubelet', 'used_percent': 48.0}]},
            },
        },
        'workload_analysis': {
            'tenant_quota_summary': {
                'scope_total': 10,
                'scope_covered': 8,
                'rows': [
                    {'namespace': 'payments', 'status': 'Covered'},
                    {'namespace': 'core', 'status': 'Covered'},
                ],
            },
        },
        'resource_forecast': {
            'forecast_payload': {
                'items': [
                    {
                        'key': 'memory',
                        'current_percent': 61.0,
                        'projected_percent': 74.0,
                        'historical_points': [{'value': 55.0}, {'value': 61.0}],
                        'forecast_points': [{'value': 68.0}, {'value': 74.0}],
                    },
                    {'key': 'cpu', 'current_percent': 28.0, 'projected_percent': 35.0},
                ],
            },
        },
        'top_consumers': {
            'consumer_payload': {
                'rows': [
                    {
                        'resource': 'Memory',
                        'namespace': 'monitoring',
                        'name': 'prometheus-0',
                        'value': 14.2,
                        'unit': 'GiB',
                    },
                    {'resource': 'Memory', 'namespace': 'logging', 'name': 'loki-0', 'value': 8.7, 'unit': 'GiB'},
                    {'resource': 'CPU', 'namespace': 'payments', 'name': 'api-0', 'value': 0.8, 'unit': 'cores'},
                ],
            },
        },
        'alerts_anomalies': {
            'summary': {
                'active_alerts': 4,
                'critical_alerts': 1,
                'warning_alerts': 2,
                'info_alerts': 1,
                'detected_anomalies': 1,
            },
            'alerts': [
                {'name': 'HighMemory', 'severity': 'critical', 'status': 'firing'},
                {'name': 'RestartBurst', 'severity': 'warning', 'status': 'firing'},
            ],
        },
        'object_storage_analysis': {
            'summary': {
                'capacity': '120 GiB',
                'usage': '76 GiB',
                'headroom': '44 GiB',
                'utilization': '63.3%',
            },
            'capacity_usage': {
                'capacity_total': '120 GiB',
                'usage_total': '76 GiB',
                'capacity_free': '44 GiB',
                'utilization': '63.3%',
            },
            'top_consumers': [
                {'name': 'thanos', 'usage_gib': 64.25, 'object_count': 9000},
                {'name': 'crm', 'usage_gib': 7.37, 'object_count': 2100},
            ],
        },
    }
    return payloads[report_type]


def test_analytics_overview_workbook_renders_two_series_comparison_without_fake_point_legend(tmp_path: Path) -> None:
    output_path = tmp_path / 'analytics-overview.xlsx'
    build_analytics_workbook(
        report_type='analytics_overview',
        payload=_representative_payload('analytics_overview'),
        context=_context('analytics_overview'),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    overview = workbook['Overview']
    chart = overview._charts[0]
    assert chart.__class__.__name__ == 'BarChart'
    assert chart.grouping == 'clustered'
    assert chart.legend is not None
    assert chart.legend.position == 'b'
    assert [series.tx.v for series in chart.series] == ['Current', 'Projected']
    assert chart.dLbls is not None
    assert chart.dLbls.showVal is True
    assert chart.dLbls.showLegendKey is False
    assert chart.dLbls.showCatName is False
    assert chart.dLbls.showSerName is False
    assert chart.dLbls.showPercent is False
    assert chart.dLbls.showBubbleSize is False
    assert chart.dLbls.dLblPos == 'outEnd'
    assert chart.x_axis.tickLblPos == 'nextTo'
    assert chart.x_axis.delete is False
    assert chart.y_axis.title is None
    assert chart.y_axis.scaling.min == 0
    assert chart.y_axis.scaling.max == 100
    assert chart.y_axis.majorUnit == 25
    assert chart.y_axis.majorGridlines is None
    assert overview.column_dimensions['J'].width == 20
    assert overview.column_dimensions['K'].width == 14
    assert overview.column_dimensions['L'].width == 14
    assert overview.column_dimensions['M'].width == 17
    assert overview.column_dimensions['N'].width == 20
    assert overview['J15'].value == 'Resource'
    assert overview['K15'].value == 'Current'
    assert overview['L15'].value == 'Projected'
    assert overview['M15'].value == 'Change (pp)'
    assert overview['N15'].value == 'Trend'
    assert overview['M16'].alignment.horizontal == 'center'
    assert overview['J16'].value == 'CPU'
    assert overview['J17'].value == 'Memory'
    assert overview['K17'].value == 64.0
    assert overview['L17'].value == 72.0


@pytest.mark.parametrize(
    ('report_type', 'chart_class'),
    [
        ('analytics_overview', 'BarChart'),
        ('resource_utilization', 'LineChart'),
        ('capacity_analysis', 'BarChart'),
        ('storage_analysis', 'BarChart'),
        ('workload_analysis', 'BarChart'),
        ('resource_forecast', 'LineChart'),
        ('top_consumers', 'BarChart'),
        ('alerts_anomalies', 'BarChart'),
        ('object_storage_analysis', 'BarChart'),
    ],
)
def test_every_analytics_category_builds_its_own_runtime_valid_workbook(
    tmp_path: Path,
    report_type: str,
    chart_class: str,
) -> None:
    output_path = tmp_path / f'{report_type}.xlsx'
    context = _context(report_type)
    if report_type == 'object_storage_analysis':
        context = replace(
            context,
            scope_label='Object Storage System',
            scope_value='minio-pgn-development',
            namespace_scope='Not Applicable',
        )

    build_analytics_workbook(
        report_type=report_type,
        payload=_representative_payload(report_type),
        context=context,
        output_path=output_path,
    )
    validate_analytics_xlsx_artifact(output_path)

    workbook = load_workbook(output_path)
    assert workbook.sheetnames[0] == 'Overview'
    assert workbook.sheetnames[-2:] == ['Metadata', 'Field Dictionary']
    assert len(workbook['Overview']._charts) == 1
    assert workbook['Overview']._charts[0].__class__.__name__ == chart_class
    metadata = workbook['Metadata']
    metadata_values = {
        metadata.cell(row=row, column=1).value: metadata.cell(row=row, column=2).value
        for row in range(1, metadata.max_row + 1)
    }
    assert metadata_values['Workbook Standard'] == ANALYTICS_WORKBOOK_STANDARD


def test_ranked_bar_normalizes_provider_native_memory_bytes_to_committed_display_unit(tmp_path: Path) -> None:
    payload = {
        'consumer_payload': {
            'rows': [
                {
                    'resource': 'Memory',
                    'namespace': 'openebs',
                    'name': 'cstor-disk-dfgm-7f4dfcd7cd',
                    'current_usage': 18_638_389_248,
                    'current_display': '17.36 GiB',
                },
                {
                    'resource': 'Memory',
                    'namespace': 'openebs',
                    'name': 'cstor-disk-thkc-55cf8fb6fb',
                    'current_usage': 18_048_593_920,
                    'current_display': '16.81 GiB',
                },
            ]
        }
    }

    prepared = prepare_analytics_workbook_data('top_consumers', payload, _context('top_consumers'))
    assert prepared.primary_chart is not None
    assert prepared.primary_chart.series[0].unit == 'GiB'
    assert prepared.primary_chart.rows[0]['value'] == pytest.approx(17.36)
    assert prepared.primary_chart.rows[1]['value'] == pytest.approx(16.81)
    assert prepared.primary_chart.total_value == pytest.approx(34.17)

    output_path = tmp_path / 'normalized-memory-ranked.xlsx'
    build_analytics_workbook(
        report_type='top_consumers',
        payload=payload,
        context=_context('top_consumers'),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    overview = workbook['Overview']
    chart = overview._charts[0]
    assert overview['L16'].value == pytest.approx(17.36)
    assert overview['L16'].number_format == '#,##0.00" GiB"'
    assert chart.y_axis.scaling.max is not None
    assert chart.y_axis.scaling.max >= 20


def test_ranked_bar_uses_compact_chart_identity_and_value_label_headroom(tmp_path: Path) -> None:
    payload = {
        'consumer_payload': {
            'rows': [
                {
                    'resource': 'Memory',
                    'namespace': 'monitoring',
                    'name': 'prometheus-prometheus-kube-prometheus-prometheus-very-long-identity',
                    'current_usage': 18_638_389_248,
                    'current_display': '17.36 GiB',
                },
                {
                    'resource': 'Memory',
                    'namespace': 'openebs',
                    'name': 'cstor-disk-dfgm-7f4dfcd7cd',
                    'current_usage': 2_147_483_648,
                    'current_display': '2.00 GiB',
                },
            ]
        }
    }
    output_path = tmp_path / 'ranked-bar-layout.xlsx'
    build_analytics_workbook(
        report_type='top_consumers',
        payload=payload,
        context=_context('top_consumers'),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    overview = workbook['Overview']
    chart = overview._charts[0]

    assert overview['K16'].value == 'prometheus-prometheus-kube-prometheus-prometheus-very-long-identity'
    assert overview['O16'].value.endswith('…')
    assert len(overview['O16'].value) <= 28
    assert chart.y_axis.scaling.min == 0
    assert chart.y_axis.scaling.max is not None
    assert chart.y_axis.scaling.max >= 20
    assert chart.legend is None
    assert chart.dLbls is not None
    assert chart.dLbls.dLblPos == 'outEnd'


def test_storage_capacity_visual_uses_one_persistent_claim_posture_not_used_and_free_as_two_resources(
    tmp_path: Path,
) -> None:
    payload = {
        'storage_summary': {
            'persistent_storage': {
                'pv': {'count': 18, 'capacity': '120 GiB'},
                'pvc': {
                    'count': 21,
                    'used_percent': 63.0,
                    'free_percent': 37.0,
                    'used': '75.6 GiB',
                    'free': '44.4 GiB',
                },
            },
            'node_filesystem': {'rows': [{'name': '/var/lib/kubelet', 'used_percent': 48.0}]},
        }
    }
    prepared = prepare_analytics_workbook_data('storage_analysis', payload, _context('storage_analysis'))
    assert prepared.primary_chart is not None
    assert prepared.primary_chart.rows == ({'category': 'Persistent Claims', 'value': 63.0},)

    output_path = tmp_path / 'storage-capacity-posture.xlsx'
    build_analytics_workbook(
        report_type='storage_analysis',
        payload=payload,
        context=_context('storage_analysis'),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    overview = workbook['Overview']
    assert overview['J16'].value == 'Persistent Claims'
    assert overview['K16'].value == pytest.approx(63.0)
    assert overview['L16'].value == pytest.approx(37.0)
    assert overview['J17'].value is None


def test_top_consumers_evidence_sheet_uses_one_normalized_numeric_value_column() -> None:
    payload = {
        'consumer_payload': {
            'rows': [
                {
                    'resource': 'Memory',
                    'namespace': 'openebs',
                    'name': 'cstor-disk-dfgm-7f4dfcd7cd',
                    'current_usage': 18_638_389_248,
                    'current_display': '17.36 GiB',
                }
            ]
        }
    }
    prepared = prepare_analytics_workbook_data('top_consumers', payload, _context('top_consumers'))
    table = next(table for table in prepared.tables if table.sheet_name == 'Top Consumers')
    assert [column.label for column in table.columns] == [
        'Resource',
        'Namespace',
        'Consumer',
        'Value',
        'Unit',
        'Status',
    ]
    assert table.rows[0]['value'] == pytest.approx(17.36)
    assert table.rows[0]['unit'] == 'GiB'
    assert 'display_value' not in {column.key for column in table.columns}


def test_resource_utilization_consumer_evidence_normalizes_memory_bytes_to_gib() -> None:
    payload = {
        'summary': {
            'cpu_utilization': 18.2,
            'memory_utilization': 61.4,
            'storage_utilization': 42.5,
            'pod_utilization': 53.0,
        },
        'top_cpu_consumers': [{'namespace': 'payments', 'pod': 'api-0', 'value': 0.71}],
        'top_memory_consumers': [{'namespace': 'openebs', 'pod': 'cstor-0', 'value': 18_638_389_248}],
    }
    prepared = prepare_analytics_workbook_data('resource_utilization', payload, _context('resource_utilization'))
    table = next(table for table in prepared.tables if table.sheet_name == 'Top Consumers')
    memory = next(row for row in table.rows if row['resource'] == 'Memory')
    assert memory['value'] == pytest.approx(18_638_389_248 / 1024**3)
    assert memory['unit'] == 'GiB'


def test_missing_evidence_status_uses_neutral_not_healthy_semantic_style(tmp_path: Path) -> None:
    payload = {
        'consumer_payload': {
            'rows': [
                {
                    'resource': 'Memory',
                    'namespace': 'openebs',
                    'name': 'cstor-0',
                    'current_usage': 2_147_483_648,
                    'current_display': '2.00 GiB',
                    'status': 'Not Available',
                }
            ]
        }
    }
    output_path = tmp_path / 'neutral-missing-evidence.xlsx'
    build_analytics_workbook(
        report_type='top_consumers',
        payload=payload,
        context=_context('top_consumers'),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    status_cell = workbook['Top Consumers']['F6']
    assert status_cell.value == 'Not Available'
    assert status_cell.fill.fgColor.rgb.endswith(KAORI_SLATE_FILL)
