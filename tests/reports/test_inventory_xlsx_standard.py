from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.chart import BarChart, DoughnutChart
from openpyxl.utils import get_column_letter

from app.reports.export.inventory_common import InventoryExportContext
from app.reports.export.tabular import TABULAR_EXPORT_STANDARD, TABULAR_SCHEMA_VERSION
from app.reports.export.xlsx.builder import (
    DATA_START_ROW,
    HEADER_ROW,
    _inventory_print_width,
    build_inventory_workbook,
    validate_inventory_xlsx_artifact,
)
from app.reports.export.xlsx.standard import (
    INVENTORY_DATA_SHEET,
    INVENTORY_DICTIONARY_SHEET,
    INVENTORY_METADATA_SHEET,
    INVENTORY_OVERVIEW_SHEET,
    INVENTORY_WORKBOOK_PROFILES,
    INVENTORY_WORKBOOK_STANDARD,
    XLSX_REPORTING_STANDARD,
)

EXPECTED_REPORT_TYPES = {
    'inventory_overview',
    'node_inventory',
    'namespace_inventory',
    'workload_inventory',
    'pod_inventory',
    'service_inventory',
    'ingress_inventory',
    'configuration_inventory',
    'storage_inventory',
    'object_storage_inventory',
    'tls_inventory',
}


def _sample_value(data_type: str, key: str) -> object:
    if data_type == 'datetime':
        return datetime(2026, 8, 3, 7, 47, tzinfo=UTC)
    if data_type == 'date':
        return datetime(2026, 8, 3, tzinfo=UTC).date()
    if data_type in {'integer', 'number', 'decimal'}:
        return 1
    if data_type == 'boolean':
        return True
    if key in {'status', 'inventory_status'}:
        return 'Current'
    if key == 'ready':
        return 'Ready'
    if key == 'status_label':
        return 'Valid'
    if key == 'usage_label':
        return 'In Use'
    if key == 'coverage_label':
        return 'Fully Covered'
    if key == 'parse_state':
        return 'Passed'
    return f'Sample {key.replace("_", " ")}'


def _payload(report_type: str) -> dict[str, object]:
    profile = INVENTORY_WORKBOOK_PROFILES[report_type]
    row: dict[str, object] = {}
    for column in profile.columns:
        source_key = column.source_key or column.key
        if column.scale_divisor:
            row[source_key] = column.scale_divisor * 2.5
        else:
            row[source_key] = _sample_value(column.data_type, source_key)
    return {
        profile.row_source: [row],
        'data_cutoff': datetime(2026, 8, 3, 7, 57, tzinfo=UTC),
        'collected_at': datetime(2026, 8, 3, 7, 47, tzinfo=UTC),
    }


def _context(report_type: str) -> InventoryExportContext:
    return InventoryExportContext(
        report_id='KAORI-INV-NOD-20260803-ABCDEF12',
        report_name=f'{report_type} validation',
        report_type=report_type,
        generated_at=datetime(2026, 8, 3, 8, 0, tzinfo=UTC),
        generated_by='Platform Operations',
        cluster_scope='K8s PGN Staging',
        namespace_scope='All Namespaces',
        period_label='Current Inventory Snapshot',
        data_cutoff=datetime(2026, 8, 3, 7, 57, tzinfo=UTC),
    )


def test_inventory_workbook_profiles_cover_every_supported_inventory_report() -> None:
    assert set(INVENTORY_WORKBOOK_PROFILES) == EXPECTED_REPORT_TYPES


def test_every_inventory_profile_declares_a_global_overview_chart() -> None:
    for report_type, profile in INVENTORY_WORKBOOK_PROFILES.items():
        assert profile.overview_chart is not None, report_type

        column_keys = {column.key for column in profile.columns}
        chart = profile.overview_chart
        assert chart.category_key in column_keys, report_type
        if chart.aggregation == 'value':
            assert chart.value_key in column_keys, report_type
        else:
            assert chart.aggregation == 'count', report_type
            assert chart.value_key is None, report_type


def test_identity_columns_are_leading_and_reusable_across_csv_xlsx() -> None:
    for report_type, profile in INVENTORY_WORKBOOK_PROFILES.items():
        keys = tuple(column.key for column in profile.columns)
        assert keys[: len(profile.identity_columns)] == profile.identity_columns, report_type


def test_status_semantics_are_owned_by_columns() -> None:
    node_profile = INVENTORY_WORKBOOK_PROFILES['node_inventory']
    semantics = {column.key: column.semantic_type for column in node_profile.columns}
    assert semantics['ready'] == 'readiness'
    assert semantics['inventory_status'] == 'inventory_state'
    assert semantics['scheduling_state'] == 'scheduling'
    assert semantics['memory_pressure'] == 'negative_condition'
    assert semantics['disk_pressure'] == 'negative_condition'


def test_all_inventory_workbooks_follow_the_reusable_four_sheet_contract(tmp_path: Path) -> None:
    for report_type, profile in INVENTORY_WORKBOOK_PROFILES.items():
        output_path = tmp_path / f'{report_type}.xlsx'
        build_inventory_workbook(
            report_type=report_type,
            payload=_payload(report_type),
            context=_context(report_type),
            output_path=output_path,
        )

        workbook = load_workbook(output_path)
        assert workbook.sheetnames == [
            INVENTORY_OVERVIEW_SHEET,
            INVENTORY_DATA_SHEET,
            INVENTORY_METADATA_SHEET,
            INVENTORY_DICTIONARY_SHEET,
        ]

        summary = workbook[INVENTORY_OVERVIEW_SHEET]
        assert summary['A1'].value == f'KAORI — {profile.title} Overview'
        assert summary['A5'].value == 'Executive Snapshot'
        assert summary['A6'].value == profile.overview_metrics[0].label
        assert summary['A7'].value == 1
        assert summary['A9'].value == 'Executive Readout'
        assert summary['A13'].value == 'Primary Visual'
        assert summary['I13'].value == 'Supporting Data'
        assert summary['A31'].value == 'Workbook Navigation'
        assert summary['A32'].value == 'Open Inventory →'
        assert summary['A32'].hyperlink.target == "#'Inventory'!A1"
        assert summary['E32'].value == 'Open Metadata →'
        assert summary['E32'].hyperlink.target == "#'Metadata'!A1"
        assert summary['I32'].value == 'Open Field Dictionary →'
        assert summary['I32'].hyperlink.target == "#'Field Dictionary'!A1"
        assert len(summary._charts) == 1

        inventory = workbook[INVENTORY_DATA_SHEET]
        expected_freeze_column = get_column_letter(max(1, len(profile.identity_columns) + 1))
        assert inventory.freeze_panes == f'{expected_freeze_column}{DATA_START_ROW}'
        assert inventory.auto_filter.ref is not None
        assert len(inventory.tables) == 1
        assert inventory['A1'].value == f'KAORI — {profile.title}'
        assert inventory['A3'].value.startswith('Reference: KAORI-INV-')
        assert inventory.cell(row=HEADER_ROW, column=1).value == profile.columns[0].label
        assert inventory.print_title_rows == f'$1:${HEADER_ROW}'
        assert inventory.page_setup.fitToWidth == _inventory_print_width(len(profile.columns))

        for cell in inventory[HEADER_ROW][: len(profile.columns)]:
            assert cell.alignment.horizontal == 'left'
            assert cell.alignment.vertical == 'center'
            assert cell.alignment.wrap_text is True
        for index, column in enumerate(profile.columns, start=1):
            cell = inventory.cell(row=DATA_START_ROW, column=index)
            expected_horizontal = 'right' if column.data_type in {'integer', 'number', 'decimal'} else 'left'
            assert cell.alignment.horizontal == expected_horizontal
            assert cell.alignment.vertical == 'center'
            assert cell.alignment.wrap_text is True

        metadata = workbook[INVENTORY_METADATA_SHEET]
        metadata_values = {
            metadata.cell(row=row, column=1).value: metadata.cell(row=row, column=2).value
            for row in range(5, metadata.max_row + 1)
        }
        assert metadata_values['Workbook Standard'] == INVENTORY_WORKBOOK_STANDARD
        assert metadata_values['XLSX Standard'] == XLSX_REPORTING_STANDARD
        assert metadata_values['Tabular Standard'] == TABULAR_EXPORT_STANDARD
        assert metadata_values['Schema Version'] == TABULAR_SCHEMA_VERSION
        assert metadata_values['Primary Format'] == 'XLSX'
        assert metadata_values['CSV Parity'] == 'Enabled'

        dictionary = workbook[INVENTORY_DICTIONARY_SHEET]
        assert dictionary.max_row == len(profile.columns) + 4
        assert dictionary.max_column == 9
        assert len(dictionary.tables) == 1


def test_distribution_overview_chart_aggregates_records_by_category(tmp_path: Path) -> None:
    report_type = 'node_inventory'
    profile = INVENTORY_WORKBOOK_PROFILES[report_type]
    output_path = tmp_path / 'nodes-overview-distribution.xlsx'
    payload = _payload(report_type)
    source_rows = payload[profile.row_source]
    assert isinstance(source_rows, list)
    base_row = dict(source_rows[0])
    rows = []
    for index, readiness in enumerate(('Ready', 'Not Ready', 'Ready'), start=1):
        row = dict(base_row)
        row['node'] = f'node-{index}'
        row['ready'] = readiness
        rows.append(row)
    payload[profile.row_source] = rows

    build_inventory_workbook(
        report_type=report_type,
        payload=payload,
        context=_context(report_type),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    summary = workbook[INVENTORY_OVERVIEW_SHEET]
    assert summary['J15'].value == 'Ready'
    assert summary['K15'].value == 'Records'
    assert summary['J16'].value == 'Ready'
    assert summary['K16'].value == 2
    assert summary['J17'].value == 'Not Ready'
    assert summary['K17'].value == 1
    assert summary['K16'].number_format == '#,##0'
    assert summary['L16'].value == pytest.approx(66.6666666667)
    assert len(summary._charts) == 1



def test_runtime_xlsx_artifact_validator_accepts_the_global_contract(tmp_path: Path) -> None:
    report_type = 'node_inventory'
    output_path = tmp_path / 'nodes-runtime-contract.xlsx'
    build_inventory_workbook(
        report_type=report_type,
        payload=_payload(report_type),
        context=_context(report_type),
        output_path=output_path,
    )

    validate_inventory_xlsx_artifact(output_path)


def test_runtime_xlsx_artifact_validator_rejects_missing_overview(tmp_path: Path) -> None:
    report_type = 'node_inventory'
    output_path = tmp_path / 'nodes-missing-overview.xlsx'
    build_inventory_workbook(
        report_type=report_type,
        payload=_payload(report_type),
        context=_context(report_type),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    workbook.remove(workbook[INVENTORY_OVERVIEW_SHEET])
    workbook.save(output_path)

    with pytest.raises(ValueError, match='Overview'):
        validate_inventory_xlsx_artifact(output_path)


def test_runtime_xlsx_artifact_validator_rejects_missing_visual_summary(tmp_path: Path) -> None:
    report_type = 'node_inventory'
    output_path = tmp_path / 'nodes-missing-chart.xlsx'
    build_inventory_workbook(
        report_type=report_type,
        payload=_payload(report_type),
        context=_context(report_type),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    workbook[INVENTORY_OVERVIEW_SHEET]._charts.clear()
    workbook.save(output_path)

    with pytest.raises(ValueError, match='chart'):
        validate_inventory_xlsx_artifact(output_path)


def test_runtime_xlsx_artifact_validator_rejects_stale_workbook_standard(tmp_path: Path) -> None:
    report_type = 'node_inventory'
    output_path = tmp_path / 'nodes-stale-standard.xlsx'
    build_inventory_workbook(
        report_type=report_type,
        payload=_payload(report_type),
        context=_context(report_type),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    metadata = workbook[INVENTORY_METADATA_SHEET]
    for row in range(5, metadata.max_row + 1):
        if metadata.cell(row=row, column=1).value == 'Workbook Standard':
            metadata.cell(row=row, column=2).value = 'kaori-inventory-xlsx-v2.0'
            break
    workbook.save(output_path)

    with pytest.raises(ValueError, match='Workbook Standard'):
        validate_inventory_xlsx_artifact(output_path)


def test_object_storage_workbook_uses_readable_gib_columns(tmp_path: Path) -> None:
    report_type = 'object_storage_inventory'
    profile = INVENTORY_WORKBOOK_PROFILES[report_type]
    output_path = tmp_path / 'object-storage.xlsx'
    payload = _payload(report_type)

    build_inventory_workbook(
        report_type=report_type,
        payload=payload,
        context=_context(report_type),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    inventory = workbook[INVENTORY_DATA_SHEET]
    headers = [inventory.cell(row=HEADER_ROW, column=index).value for index in range(1, len(profile.columns) + 1)]
    assert 'Usage GiB' in headers
    assert 'Quota GiB' in headers
    assert 'Usage Bytes' not in headers
    assert 'Quota Bytes' not in headers
    usage_index = headers.index('Usage GiB') + 1
    quota_index = headers.index('Quota GiB') + 1
    assert inventory.cell(row=DATA_START_ROW, column=usage_index).value == 2.5
    assert inventory.cell(row=DATA_START_ROW, column=quota_index).value == 2.5
    utilization_index = headers.index('Quota Utilization %') + 1
    assert inventory.cell(row=DATA_START_ROW, column=utilization_index).value == 100.0

def test_numeric_missing_values_stay_blank_and_text_missing_values_are_explicit(tmp_path: Path) -> None:
    report_type = 'node_inventory'
    profile = INVENTORY_WORKBOOK_PROFILES[report_type]
    output_path = tmp_path / 'nodes-empty-fields.xlsx'
    row = {column.key: None for column in profile.columns}

    build_inventory_workbook(
        report_type=report_type,
        payload={profile.row_source: [row]},
        context=_context(report_type),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    inventory = workbook[INVENTORY_DATA_SHEET]
    for index, column in enumerate(profile.columns, start=1):
        value = inventory.cell(row=DATA_START_ROW, column=index).value
        if column.data_type in {'integer', 'number', 'decimal'}:
            assert value is None
        else:
            assert value == 'Not Available'


def test_tls_inventory_never_exports_private_key_or_secret_values(tmp_path: Path) -> None:
    profile = INVENTORY_WORKBOOK_PROFILES['tls_inventory']
    labels = {column.label.lower() for column in profile.columns}
    keys = {column.key.lower() for column in profile.columns}
    forbidden = {'private key', 'tls.key', 'secret value', 'secret data', 'raw secret'}

    assert not labels.intersection(forbidden)
    assert not keys.intersection({'private_key', 'tls_key', 'secret_value', 'secret_data'})

    output_path = tmp_path / 'tls-inventory.xlsx'
    build_inventory_workbook(
        report_type='tls_inventory',
        payload=_payload('tls_inventory'),
        context=_context('tls_inventory'),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    metadata = workbook[INVENTORY_METADATA_SHEET]
    metadata_values = {
        metadata.cell(row=row, column=1).value: metadata.cell(row=row, column=2).value
        for row in range(5, metadata.max_row + 1)
    }
    assert metadata_values['Private Key Export'] == 'Excluded'


def test_empty_inventory_workbook_keeps_a_valid_unmerged_excel_table(tmp_path: Path) -> None:
    output_path = tmp_path / 'empty-services.xlsx'
    build_inventory_workbook(
        report_type='service_inventory',
        payload={'inventory_assets': []},
        context=_context('service_inventory'),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    inventory = workbook[INVENTORY_DATA_SHEET]
    assert inventory.cell(row=DATA_START_ROW, column=1).value == (
        'No committed inventory rows were found for the selected reporting boundary.'
    )
    assert not any(str(DATA_START_ROW) in str(cell_range) for cell_range in inventory.merged_cells.ranges)
    assert len(inventory.tables) == 1



def test_global_xlsx_v2_profiles_declare_domain_metrics_and_chart_semantics() -> None:
    assert XLSX_REPORTING_STANDARD == 'kaori-reporting-xlsx-v2.0'
    assert INVENTORY_WORKBOOK_STANDARD == 'kaori-inventory-xlsx-v4.0'

    for report_type, profile in INVENTORY_WORKBOOK_PROFILES.items():
        assert len(profile.overview_metrics) == 4, report_type
        assert profile.record_label, report_type
        assert profile.overview_chart.visual_type in {'ranked_bar', 'doughnut'}, report_type
        assert profile.overview_chart.top_n <= 8, report_type
        metric_labels = {metric.label for metric in profile.overview_metrics}
        assert 'Columns' not in metric_labels, report_type
        assert 'Identity Fields' not in metric_labels, report_type


def test_object_storage_overview_is_executive_and_ranked_bar_is_directly_labeled(tmp_path: Path) -> None:
    report_type = 'object_storage_inventory'
    profile = INVENTORY_WORKBOOK_PROFILES[report_type]
    output_path = tmp_path / 'object-storage-executive.xlsx'
    base_payload = _payload(report_type)
    source_rows = base_payload[profile.row_source]
    assert isinstance(source_rows, list)
    base_row = dict(source_rows[0])
    gib = 1073741824.0
    rows = []
    for name, usage_gib, objects, quota_gib in (
        ('thanos', 64.25, 1200, 100.0),
        ('shg-crm', 7.37, 350, 20.0),
        ('pgnmobile', 4.61, 180, 10.0),
    ):
        row = dict(base_row)
        row['name'] = name
        row['size_bytes'] = usage_gib * gib
        row['object_count'] = objects
        row['quota_bytes'] = quota_gib * gib
        rows.append(row)
    base_payload[profile.row_source] = rows

    build_inventory_workbook(
        report_type=report_type,
        payload=base_payload,
        context=_context(report_type),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    overview = workbook[INVENTORY_OVERVIEW_SHEET]
    assert overview['A5'].value == 'Executive Snapshot'
    assert overview['A6'].value == 'Buckets / Containers'
    assert overview['A7'].value == 3
    assert overview['D6'].value == 'Total Usage GiB'
    assert overview['D7'].value == pytest.approx(76.23)
    assert overview['G6'].value == 'Total Objects'
    assert overview['G7'].value == 1730
    assert overview['J6'].value == 'Quota Coverage'
    assert overview['J7'].value == pytest.approx(100.0)
    assert overview['A9'].value == 'Executive Readout'
    assert 'thanos' in str(overview['A10'].value)
    assert '64.25 GiB' in str(overview['A10'].value)
    assert overview['A13'].value == 'Primary Visual'
    assert overview['I13'].value == 'Supporting Data'
    assert [overview.cell(row=15, column=column).value for column in range(9, 13)] == [
        'Rank', 'Bucket / Container', 'Usage GiB', 'Share of Total'
    ]
    assert overview['J16'].value == 'thanos'
    assert overview['K16'].value == pytest.approx(64.25)
    assert overview['L16'].value == pytest.approx(64.25 / 76.23 * 100)

    assert len(overview._charts) == 1
    chart = overview._charts[0]
    assert isinstance(chart, BarChart)
    assert chart.type == 'bar'
    assert chart.varyColors is False
    assert chart.legend is None
    assert chart.dLbls is not None and chart.dLbls.showVal is True
    assert chart.dLbls.numFmt == '#,##0.00" GiB"'
    assert chart.y_axis.scaling.min == 0
    assert chart.x_axis.scaling.orientation == 'maxMin'


def test_ranked_bar_supporting_share_uses_full_dataset_denominator_not_only_top_n(tmp_path: Path) -> None:
    report_type = 'object_storage_inventory'
    profile = INVENTORY_WORKBOOK_PROFILES[report_type]
    output_path = tmp_path / 'object-storage-top-n-share.xlsx'
    payload = _payload(report_type)
    source_rows = payload[profile.row_source]
    assert isinstance(source_rows, list)
    base_row = dict(source_rows[0])
    gib = 1073741824.0
    rows = []
    for index, usage_gib in enumerate(range(90, 0, -10), 1):
        row = dict(base_row)
        row['name'] = f'bucket-{index}'
        row['size_bytes'] = usage_gib * gib
        row['object_count'] = index
        row['quota_bytes'] = 100 * gib
        rows.append(row)
    payload[profile.row_source] = rows

    build_inventory_workbook(
        report_type=report_type,
        payload=payload,
        context=_context(report_type),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    overview = workbook[INVENTORY_OVERVIEW_SHEET]
    assert overview['J16'].value == 'bucket-1'
    assert overview['K16'].value == pytest.approx(90.0)
    assert overview['L16'].value == pytest.approx(90.0 / 450.0 * 100.0)
    assert '20.0% of total Usage GiB' in str(overview['A10'].value)


def test_distribution_chart_uses_doughnut_legend_percent_labels_and_domain_kpis(tmp_path: Path) -> None:
    report_type = 'node_inventory'
    profile = INVENTORY_WORKBOOK_PROFILES[report_type]
    output_path = tmp_path / 'nodes-executive.xlsx'
    payload = _payload(report_type)
    source_rows = payload[profile.row_source]
    assert isinstance(source_rows, list)
    base_row = dict(source_rows[0])
    rows = []
    for index, (readiness, role) in enumerate((('Ready', 'worker'), ('Ready', 'worker'), ('Not Ready', 'control-plane')), 1):
        row = dict(base_row)
        row['node'] = f'node-{index}'
        row['ready'] = readiness
        row['role'] = role
        rows.append(row)
    payload[profile.row_source] = rows

    build_inventory_workbook(
        report_type=report_type,
        payload=payload,
        context=_context(report_type),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    overview = workbook[INVENTORY_OVERVIEW_SHEET]
    assert overview['A6'].value == 'Total Nodes'
    assert overview['A7'].value == 3
    assert overview['D6'].value == 'Ready'
    assert overview['D7'].value == 2
    assert overview['G6'].value == 'Not Ready'
    assert overview['G7'].value == 1
    assert overview['J6'].value == 'Worker Nodes'
    assert overview['J7'].value == 2
    assert overview['J16'].value == 'Ready'
    assert overview['K16'].value == 2
    assert overview['L16'].value == pytest.approx(66.6666666667)

    chart = overview._charts[0]
    assert isinstance(chart, DoughnutChart)
    assert chart.legend is not None
    assert chart.legend.position == 'b'
    assert chart.dLbls is not None and chart.dLbls.showPercent is True
    assert chart.dLbls.showLegendKey is False
    assert chart.dLbls.showVal is False
    assert chart.dLbls.showCatName is False
    assert chart.dLbls.showSerName is False
    assert chart.dLbls.showLeaderLines is True
    assert chart.dLbls.dLblPos == 'bestFit'
    assert chart.holeSize == 62


def test_doughnut_groups_categories_beyond_declared_limit_into_other(tmp_path: Path) -> None:
    report_type = 'node_inventory'
    profile = INVENTORY_WORKBOOK_PROFILES[report_type]
    output_path = tmp_path / 'nodes-doughnut-other.xlsx'
    payload = _payload(report_type)
    source_rows = payload[profile.row_source]
    assert isinstance(source_rows, list)
    base_row = dict(source_rows[0])
    rows = []
    categories = ('Ready', 'Not Ready', 'Unknown A', 'Unknown B', 'Unknown C', 'Unknown D', 'Unknown E')
    for index, readiness in enumerate(categories, 1):
        row = dict(base_row)
        row['node'] = f'node-{index}'
        row['ready'] = readiness
        row['role'] = 'worker'
        rows.append(row)
    payload[profile.row_source] = rows

    build_inventory_workbook(
        report_type=report_type,
        payload=payload,
        context=_context(report_type),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    overview = workbook[INVENTORY_OVERVIEW_SHEET]
    supporting = [
        (overview.cell(row=row, column=10).value, overview.cell(row=row, column=11).value)
        for row in range(16, 22)
    ]
    assert len([category for category, _ in supporting if category is not None]) == 6
    assert ('Other', 2) in supporting
    other_row = next(row for row in range(16, 22) if overview.cell(row=row, column=10).value == 'Other')
    assert overview.cell(row=other_row, column=12).value == pytest.approx(2 / 7 * 100)


def test_secondary_sheets_link_back_to_overview_and_numeric_inventory_is_right_aligned(tmp_path: Path) -> None:
    report_type = 'object_storage_inventory'
    output_path = tmp_path / 'object-storage-navigation.xlsx'
    build_inventory_workbook(
        report_type=report_type,
        payload=_payload(report_type),
        context=_context(report_type),
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    inventory = workbook[INVENTORY_DATA_SHEET]
    metadata = workbook[INVENTORY_METADATA_SHEET]
    dictionary = workbook[INVENTORY_DICTIONARY_SHEET]

    assert inventory['A4'].value == '← Overview'
    assert inventory['A4'].hyperlink.target == "#'Overview'!A1"
    assert metadata['A3'].value == '← Overview'
    assert metadata['A3'].hyperlink.target == "#'Overview'!A1"
    assert dictionary['A3'].value == '← Overview'
    assert dictionary['A3'].hyperlink.target == "#'Overview'!A1"
    assert inventory.sheet_view.zoomScale == 85
    assert metadata.sheet_view.zoomScale == 90
    assert dictionary.sheet_view.zoomScale == 85

    headers = [inventory.cell(row=HEADER_ROW, column=i).value for i in range(1, inventory.max_column + 1)]
    usage_col = headers.index('Usage GiB') + 1
    object_count_col = headers.index('Object Count') + 1
    assert inventory.cell(row=DATA_START_ROW, column=usage_col).alignment.horizontal == 'right'
    assert inventory.cell(row=DATA_START_ROW, column=object_count_col).alignment.horizontal == 'right'
    assert inventory.cell(row=DATA_START_ROW, column=1).alignment.horizontal == 'left'

    section_labels = {metadata.cell(row=row, column=1).value for row in range(5, metadata.max_row + 1)}
    assert {'REPORT IDENTITY', 'REPORTING BOUNDARY', 'WORKBOOK CONTRACT', 'DATA HANDLING', 'SECURITY'} <= section_labels

    for sheet in (inventory, metadata, dictionary):
        assert all(table.tableStyleInfo.name == 'TableStyleLight1' for table in sheet.tables.values())


def test_runtime_validator_rejects_missing_executive_readout(tmp_path: Path) -> None:
    output_path = tmp_path / 'nodes-missing-readout.xlsx'
    build_inventory_workbook(
        report_type='node_inventory',
        payload=_payload('node_inventory'),
        context=_context('node_inventory'),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    workbook[INVENTORY_OVERVIEW_SHEET]['A9'] = None
    workbook.save(output_path)

    with pytest.raises(ValueError, match='Executive Readout'):
        validate_inventory_xlsx_artifact(output_path)


def test_runtime_validator_rejects_ranked_bar_vary_colors(tmp_path: Path) -> None:
    output_path = tmp_path / 'object-storage-rainbow-chart.xlsx'
    build_inventory_workbook(
        report_type='object_storage_inventory',
        payload=_payload('object_storage_inventory'),
        context=_context('object_storage_inventory'),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    overview = workbook[INVENTORY_OVERVIEW_SHEET]
    assert overview._charts
    overview._charts[0].varyColors = True
    workbook.save(output_path)

    with pytest.raises(ValueError, match='varyColors'):
        validate_inventory_xlsx_artifact(output_path)


def test_runtime_validator_rejects_ranked_bar_without_direct_value_labels(tmp_path: Path) -> None:
    output_path = tmp_path / 'object-storage-missing-direct-labels.xlsx'
    build_inventory_workbook(
        report_type='object_storage_inventory',
        payload=_payload('object_storage_inventory'),
        context=_context('object_storage_inventory'),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    overview = workbook[INVENTORY_OVERVIEW_SHEET]
    assert overview._charts and overview._charts[0].dLbls is not None
    overview._charts[0].dLbls.showVal = False
    workbook.save(output_path)

    with pytest.raises(ValueError, match='direct value labels'):
        validate_inventory_xlsx_artifact(output_path)


def test_runtime_validator_rejects_doughnut_without_legend(tmp_path: Path) -> None:
    output_path = tmp_path / 'nodes-missing-doughnut-legend.xlsx'
    build_inventory_workbook(
        report_type='node_inventory',
        payload=_payload('node_inventory'),
        context=_context('node_inventory'),
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    overview = workbook[INVENTORY_OVERVIEW_SHEET]
    assert overview._charts
    overview._charts[0].legend = None
    workbook.save(output_path)

    with pytest.raises(ValueError, match='doughnut legend'):
        validate_inventory_xlsx_artifact(output_path)
