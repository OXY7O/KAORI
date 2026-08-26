from __future__ import annotations

from collections import Counter, defaultdict
from typing import Any

from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.export.inventory_common import InventoryExportDataset
from app.reports.export.tabular import protect_spreadsheet_formula, serialize_tabular_value
from app.reports.export.xlsx.schema import InventoryOverviewMetric
from app.reports.export.xlsx.standard import INVENTORY_OVERVIEW_SHEET
from app.reports.export.xlsx.style import (
    KAORI_ACCENT,
    KAORI_AMBER_TEXT,
    KAORI_BLUE_TEXT,
    KAORI_GREEN_TEXT,
    KAORI_MUTED,
    KAORI_PRIMARY,
    KAORI_PRIMARY_DARK,
    KAORI_RED_TEXT,
    KAORI_SLATE_TEXT,
    configure_sheet_page,
    set_footer,
    style_body,
    style_context,
    style_header,
    style_navigation_link,
    style_readout,
    style_section_header,
    style_subtitle,
    style_summary_label,
    style_summary_value,
    style_title,
)

_SNAPSHOT_HEADER_ROW = 5
_KPI_LABEL_ROW = 6
_KPI_VALUE_ROW = 7
_READOUT_HEADER_ROW = 9
_READOUT_VALUE_ROW = 10
_VISUAL_HEADER_ROW = 13
_SUPPORTING_HEADER_ROW = 13
_SUPPORTING_COLUMNS = (9, 10, 11, 12)  # I:L
_SUPPORTING_TABLE_HEADER_ROW = 14
_SUPPORTING_DATA_START_ROW = 15
_NAVIGATION_HEADER_ROW = 31
_NAVIGATION_LINK_ROW = 32
_OVERVIEW_LAST_ROW = 33


def build_inventory_overview_sheet(sheet: Worksheet, dataset: InventoryExportDataset) -> None:
    _write_overview_header(sheet, dataset)
    _write_overview_kpis(sheet, dataset)
    _write_executive_readout(sheet, dataset)
    _write_overview_visual(sheet, dataset)
    _write_overview_navigation(sheet)

    for column, width in {
        'A': 13,
        'B': 13,
        'C': 13,
        'D': 13,
        'E': 13,
        'F': 13,
        'G': 13,
        'H': 13,
        'I': 7,
        'J': 30,
        'K': 15,
        'L': 12,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A5'
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_view.zoomScaleNormal = 90
    configure_sheet_page(
        sheet,
        print_area=f'A1:L{_OVERVIEW_LAST_ROW}',
        print_title_rows='$1:$3',
        orientation='landscape',
        fit_to_width=1,
    )
    set_footer(sheet, str(dataset.context.report_id))


def _write_overview_header(sheet: Worksheet, dataset: InventoryExportDataset) -> None:
    sheet.merge_cells('A1:L1')
    title = sheet['A1']
    title.value = f'KAORI — {dataset.profile.title} Overview'
    style_title(title)
    sheet.row_dimensions[1].height = 32

    sheet.merge_cells('A2:L2')
    subtitle = sheet['A2']
    subtitle.value = dataset.profile.description
    style_subtitle(subtitle)
    sheet.row_dimensions[2].height = 30

    sheet.merge_cells('A3:L3')
    context_cell = sheet['A3']
    context_cell.value = _overview_context_line(dataset)
    style_context(context_cell)
    sheet.row_dimensions[3].height = 24
    sheet.row_dimensions[4].height = 8


def _write_overview_kpis(sheet: Worksheet, dataset: InventoryExportDataset) -> None:
    sheet.merge_cells(start_row=_SNAPSHOT_HEADER_ROW, start_column=1, end_row=_SNAPSHOT_HEADER_ROW, end_column=12)
    sheet.cell(row=_SNAPSHOT_HEADER_ROW, column=1, value='Executive Snapshot')
    style_section_header(sheet.cell(row=_SNAPSHOT_HEADER_ROW, column=1))
    sheet.row_dimensions[_SNAPSHOT_HEADER_ROW].height = 22

    spans = ((1, 3), (4, 6), (7, 9), (10, 12))
    metrics = dataset.profile.overview_metrics[:4]
    for metric, (start_col, end_col) in zip(metrics, spans, strict=False):
        value = _metric_value(dataset, metric)
        _write_merged_card(
            sheet,
            row=_KPI_LABEL_ROW,
            start_col=start_col,
            end_col=end_col,
            value=metric.label,
            style=style_summary_label,
        )
        _write_merged_card(
            sheet,
            row=_KPI_VALUE_ROW,
            start_col=start_col,
            end_col=end_col,
            value=value,
            style=style_summary_value,
        )
        sheet.cell(row=_KPI_VALUE_ROW, column=start_col).number_format = metric.number_format

    sheet.row_dimensions[_KPI_LABEL_ROW].height = 22
    sheet.row_dimensions[_KPI_VALUE_ROW].height = 34
    sheet.row_dimensions[8].height = 8


def _write_executive_readout(sheet: Worksheet, dataset: InventoryExportDataset) -> None:
    sheet.merge_cells(start_row=_READOUT_HEADER_ROW, start_column=1, end_row=_READOUT_HEADER_ROW, end_column=12)
    sheet.cell(row=_READOUT_HEADER_ROW, column=1, value='Executive Readout')
    style_section_header(sheet.cell(row=_READOUT_HEADER_ROW, column=1))
    sheet.row_dimensions[_READOUT_HEADER_ROW].height = 22

    sheet.merge_cells(start_row=_READOUT_VALUE_ROW, start_column=1, end_row=_READOUT_VALUE_ROW + 1, end_column=12)
    cell = sheet.cell(row=_READOUT_VALUE_ROW, column=1, value=_executive_readout(dataset))
    style_readout(cell)
    sheet.row_dimensions[_READOUT_VALUE_ROW].height = 24
    sheet.row_dimensions[_READOUT_VALUE_ROW + 1].height = 24
    sheet.row_dimensions[12].height = 8


def _write_overview_visual(sheet: Worksheet, dataset: InventoryExportDataset) -> None:
    sheet.merge_cells(start_row=_VISUAL_HEADER_ROW, start_column=1, end_row=_VISUAL_HEADER_ROW, end_column=8)
    sheet.cell(row=_VISUAL_HEADER_ROW, column=1, value='Primary Visual')
    style_section_header(sheet.cell(row=_VISUAL_HEADER_ROW, column=1))

    sheet.merge_cells(start_row=_SUPPORTING_HEADER_ROW, start_column=9, end_row=_SUPPORTING_HEADER_ROW, end_column=12)
    sheet.cell(row=_SUPPORTING_HEADER_ROW, column=9, value='Supporting Data')
    style_section_header(sheet.cell(row=_SUPPORTING_HEADER_ROW, column=9))
    sheet.row_dimensions[_VISUAL_HEADER_ROW].height = 22

    chart_rows = _collect_chart_rows(dataset)
    if not chart_rows:
        sheet.merge_cells(start_row=14, start_column=1, end_row=16, end_column=8)
        message = sheet.cell(
            row=14,
            column=1,
            value='No committed chartable evidence is available for this workbook scope.',
        )
        style_readout(message)
        sheet.merge_cells(start_row=14, start_column=9, end_row=16, end_column=12)
        support_message = sheet.cell(row=14, column=9, value='No supporting values are available.')
        style_body(support_message)
        return

    _write_supporting_data(sheet, dataset, chart_rows)
    _add_overview_chart(sheet, dataset, chart_rows)


def _write_supporting_data(
    sheet: Worksheet,
    dataset: InventoryExportDataset,
    chart_rows: list[tuple[str, int | float]],
) -> None:
    chart = dataset.profile.overview_chart
    category_column = _profile_column(dataset, chart.category_key)
    value_column = _profile_column(dataset, chart.value_key) if chart.value_key else None
    category_label = category_column.label if category_column is not None else 'Category'
    value_label = 'Records' if chart.aggregation == 'count' else (value_column.label if value_column else 'Value')
    number_format = '#,##0' if chart.aggregation == 'count' else ((value_column.number_format if value_column else None) or '#,##0.00')

    headers = ('Rank', category_label, value_label, 'Share of Total')
    for column_index, value in zip(_SUPPORTING_COLUMNS, headers, strict=True):
        cell = sheet.cell(row=_SUPPORTING_TABLE_HEADER_ROW, column=column_index, value=value)
        style_header(cell)
    sheet.row_dimensions[_SUPPORTING_TABLE_HEADER_ROW].height = 28

    total = _chart_total(dataset)
    for rank, (category, value) in enumerate(chart_rows, start=1):
        row_index = _SUPPORTING_DATA_START_ROW + rank - 1
        share = (float(value) / total * 100.0) if total > 0 else 0.0
        values = (rank, category, value, share)
        for column_index, cell_value in zip(_SUPPORTING_COLUMNS, values, strict=True):
            cell = sheet.cell(row=row_index, column=column_index, value=cell_value)
            horizontal = 'right' if column_index in {9, 11, 12} else 'left'
            style_body(cell, alternate=rank % 2 == 0, horizontal=horizontal)
        sheet.cell(row=row_index, column=11).number_format = number_format
        sheet.cell(row=row_index, column=12).number_format = '0.0"%"'
        sheet.row_dimensions[row_index].height = 22


def _add_overview_chart(
    sheet: Worksheet,
    dataset: InventoryExportDataset,
    chart_rows: list[tuple[str, int | float]],
) -> None:
    chart = dataset.profile.overview_chart
    end_row = _SUPPORTING_DATA_START_ROW + len(chart_rows) - 1
    values = Reference(sheet, min_col=11, min_row=_SUPPORTING_TABLE_HEADER_ROW, max_row=end_row)
    categories = Reference(sheet, min_col=10, min_row=_SUPPORTING_DATA_START_ROW, max_row=end_row)

    if chart.visual_type == 'ranked_bar':
        visual = BarChart()
        visual.type = 'bar'
        visual.grouping = 'clustered'
        visual.varyColors = False
        visual.title = chart.title
        visual.height = 8.8
        visual.width = 15.8
        visual.gapWidth = 55
        visual.add_data(values, titles_from_data=True)
        visual.set_categories(categories)
        visual.legend = None
        visual.dLbls = DataLabelList(showVal=True, dLblPos='outEnd')
        value_column = _profile_column(dataset, chart.value_key) if chart.value_key else None
        number_format = '#,##0' if chart.aggregation == 'count' else ((value_column.number_format if value_column else None) or '#,##0.00')
        if chart.aggregation == 'value' and value_column is not None and value_column.unit:
            safe_unit = value_column.unit.replace('"', '')
            number_format = f'{number_format}" {safe_unit}"'
        visual.dLbls.numFmt = number_format
        visual.x_axis.scaling.orientation = 'maxMin'
        visual.x_axis.tickLblPos = 'low'
        visual.y_axis.scaling.min = 0
        visual.y_axis.majorGridlines = None
        visual.y_axis.title = 'Records' if chart.aggregation == 'count' else (value_column.label if value_column else 'Value')
        visual.series[0].graphicalProperties.solidFill = KAORI_PRIMARY
        visual.series[0].graphicalProperties.line.solidFill = KAORI_PRIMARY_DARK
        _remove_chart_frame(visual)
        sheet.add_chart(visual, 'A14')
        return

    visual = DoughnutChart()
    visual.title = chart.title
    visual.holeSize = 60
    visual.firstSliceAng = 270
    visual.varyColors = True
    visual.height = 8.6
    visual.width = 14.8
    visual.add_data(values, titles_from_data=True)
    visual.set_categories(categories)
    visual.legend.position = 'r'
    visual.dLbls = DataLabelList(showPercent=True, showLeaderLines=True)
    points: list[DataPoint] = []
    for index, (category, _) in enumerate(chart_rows):
        point = DataPoint(idx=index)
        point.graphicalProperties.solidFill = _category_color(category, index)
        points.append(point)
    visual.series[0].data_points = points
    _remove_chart_frame(visual)
    sheet.add_chart(visual, 'A14')


def _remove_chart_frame(chart: BarChart | DoughnutChart) -> None:
    chart.graphical_properties = GraphicalProperties()
    chart.graphical_properties.noFill = True
    chart.graphical_properties.line.noFill = True
    chart.graphical_properties.line.prstDash = None


def _write_overview_navigation(sheet: Worksheet) -> None:
    sheet.merge_cells(start_row=_NAVIGATION_HEADER_ROW, start_column=1, end_row=_NAVIGATION_HEADER_ROW, end_column=12)
    sheet.cell(row=_NAVIGATION_HEADER_ROW, column=1, value='Workbook Navigation')
    style_section_header(sheet.cell(row=_NAVIGATION_HEADER_ROW, column=1))
    sheet.row_dimensions[_NAVIGATION_HEADER_ROW].height = 22

    navigation = tuple(
        target_sheet for target_sheet in sheet.parent.sheetnames if target_sheet != INVENTORY_OVERVIEW_SHEET
    )
    spans = ((1, 4), (5, 8), (9, 12))
    for target_sheet, (start_col, end_col) in zip(navigation, spans, strict=False):
        for column in range(start_col, end_col + 1):
            style_navigation_link(sheet.cell(row=_NAVIGATION_LINK_ROW, column=column))
        sheet.merge_cells(
            start_row=_NAVIGATION_LINK_ROW,
            start_column=start_col,
            end_row=_NAVIGATION_LINK_ROW,
            end_column=end_col,
        )
        cell = sheet.cell(row=_NAVIGATION_LINK_ROW, column=start_col, value=f'Open {target_sheet} →')
        cell.hyperlink = f"#'{target_sheet}'!A1"
        style_navigation_link(cell)
    sheet.row_dimensions[_NAVIGATION_LINK_ROW].height = 26


def _metric_value(dataset: InventoryExportDataset, metric: InventoryOverviewMetric) -> int | float:
    if metric.aggregation == 'row_count':
        return dataset.row_count
    if metric.key is None:
        return 0
    values = [row.get(metric.key) for row in dataset.rows]
    if metric.aggregation == 'count':
        return sum(1 for value in values if _metric_matches(value, metric.match_values, metric.exclude_values))
    if metric.aggregation == 'sum':
        return sum(_numeric_value(value) for value in values)
    if metric.aggregation == 'distinct_count':
        return len({_safe_text_value(value).casefold() for value in values if value not in (None, '')})
    if metric.aggregation == 'coverage_percent':
        if not values:
            return 0.0
        covered = sum(1 for value in values if _is_positive_number(value))
        return covered / len(values) * 100.0
    if metric.aggregation == 'leading_share':
        numeric = [_numeric_value(value) for value in values if _is_number(value)]
        total = sum(numeric)
        return (max(numeric) / total * 100.0) if numeric and total > 0 else 0.0
    raise ValueError(f'Unsupported XLSX overview metric aggregation: {metric.aggregation!r}.')


def _metric_matches(
    value: Any,
    match_values: tuple[str | bool, ...],
    exclude_values: tuple[str | bool, ...],
) -> bool:
    if value in (None, ''):
        return False
    if match_values:
        return any(_same_value(value, candidate) for candidate in match_values)
    if exclude_values:
        return not any(_same_value(value, candidate) for candidate in exclude_values)
    return True


def _same_value(value: Any, candidate: str | bool) -> bool:
    if isinstance(value, str) and isinstance(candidate, str):
        return value.strip().casefold() == candidate.strip().casefold()
    return value == candidate


def _collect_chart_rows(dataset: InventoryExportDataset) -> list[tuple[str, int | float]]:
    chart = dataset.profile.overview_chart
    rows = _collect_all_chart_rows(dataset)
    if chart.visual_type != 'doughnut' or len(rows) <= chart.top_n:
        return rows[: chart.top_n]

    visible_count = chart.top_n - 1
    visible_rows = rows[:visible_count]
    other_total = sum(float(value) for _, value in rows[visible_count:])
    return [*visible_rows, ('Other', other_total)]


def _collect_all_chart_rows(dataset: InventoryExportDataset) -> list[tuple[str, int | float]]:
    chart = dataset.profile.overview_chart
    if chart.aggregation == 'count':
        counts: Counter[str] = Counter()
        for row in dataset.rows:
            category = row.get(chart.category_key)
            if category in (None, ''):
                continue
            counts[_safe_text_value(category)] += 1
        rows: list[tuple[str, int | float]] = list(counts.items())
    elif chart.aggregation == 'value' and chart.value_key is not None:
        totals: defaultdict[str, float] = defaultdict(float)
        for row in dataset.rows:
            category = row.get(chart.category_key)
            value = row.get(chart.value_key)
            if category in (None, '') or not _is_number(value):
                continue
            totals[_safe_text_value(category)] += float(value)
        rows = list(totals.items())
    else:
        raise ValueError(f'Unsupported XLSX overview chart aggregation: {chart.aggregation!r}.')

    rows.sort(key=lambda item: (-float(item[1]), item[0].casefold()))
    return rows


def _chart_total(dataset: InventoryExportDataset) -> float:
    return sum(float(value) for _, value in _collect_all_chart_rows(dataset))


def _executive_readout(dataset: InventoryExportDataset) -> str:
    if not dataset.rows:
        return 'No committed inventory records are available for the selected reporting boundary.'
    chart_rows = _collect_chart_rows(dataset)
    count_text = f'{dataset.row_count:,} {dataset.profile.record_label}'
    if not chart_rows:
        return f'{count_text} are represented in this committed snapshot; no chartable evidence is available.'

    category, value = chart_rows[0]
    total = _chart_total(dataset)
    share = (float(value) / total * 100.0) if total > 0 else 0.0
    chart = dataset.profile.overview_chart
    if chart.aggregation == 'value':
        value_column = _profile_column(dataset, chart.value_key)
        formatted_value = _format_chart_value(value, value_column.number_format if value_column else None)
        unit = f' {value_column.unit}' if value_column is not None and value_column.unit else ''
        measure_label = value_column.label if value_column is not None else 'the charted measure'
        return (
            f'{count_text} are represented in this committed snapshot. {category} is the leading consumer '
            f'at {formatted_value}{unit}, representing {share:.1f}% of total {measure_label}.'
        )
    return (
        f'{count_text} are represented in this committed snapshot. {category} is the leading category '
        f'with {int(value):,} {dataset.profile.record_label}, representing {share:.1f}% of chartable records.'
    )


def _overview_context_line(dataset: InventoryExportDataset) -> str:
    context = dataset.context
    cutoff = serialize_tabular_value(context.data_cutoff) or 'Not Available'
    return (
        f'{context.scope_label}: {_overview_value(context.cluster_scope)}  |  '
        f'Namespace: {_overview_value(context.namespace_scope)}  |  '
        f'Period: {_overview_value(context.period_label)}  |  Data Cut-off: {cutoff}'
    )


def _write_merged_card(
    sheet: Worksheet,
    *,
    row: int,
    start_col: int,
    end_col: int,
    value: Any,
    style: Any,
) -> None:
    for column in range(start_col, end_col + 1):
        style(sheet.cell(row=row, column=column))
    sheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    cell = sheet.cell(row=row, column=start_col, value=value)
    style(cell)


def _profile_column(dataset: InventoryExportDataset, key: str | None) -> Any:
    if key is None:
        return None
    return next((column for column in dataset.profile.columns if column.key == key), None)


def _category_color(category: str, index: int) -> str:
    text = category.casefold()
    if any(token in text for token in ('not ready', 'critical', 'expired', 'invalid', 'failed', 'error')):
        return KAORI_RED_TEXT
    if any(token in text for token in ('warning', 'review', 'renewal', 'pending', 'terminating', 'partial')):
        return KAORI_AMBER_TEXT
    if any(token in text for token in ('ready', 'valid', 'active', 'bound', 'running', 'current', 'enabled', 'passed')):
        return KAORI_GREEN_TEXT
    if any(token in text for token in ('unknown', 'unavailable', 'disabled', 'not available')):
        return KAORI_SLATE_TEXT
    palette = (KAORI_PRIMARY, KAORI_ACCENT, KAORI_BLUE_TEXT, KAORI_PRIMARY_DARK, KAORI_MUTED)
    return palette[index % len(palette)]


def _format_chart_value(value: int | float, number_format: str | None) -> str:
    if number_format and '.00' in number_format:
        return f'{float(value):,.2f}'
    if isinstance(value, int) or float(value).is_integer():
        return f'{int(value):,}'
    return f'{float(value):,.2f}'


def _numeric_value(value: Any) -> float:
    return float(value) if _is_number(value) else 0.0


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _is_positive_number(value: Any) -> bool:
    return _is_number(value) and float(value) > 0


def _safe_text_value(value: Any) -> str:
    return protect_spreadsheet_formula(str(value))


def _overview_value(value: Any) -> str:
    if value is None or value == '':
        return 'Not Available'
    return _safe_text_value(value)
