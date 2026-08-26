from __future__ import annotations

import re
from collections.abc import Mapping
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.export.tabular import (
    TABULAR_EXPORT_STANDARD,
    TABULAR_SCHEMA_VERSION,
    protect_spreadsheet_formula,
    serialize_tabular_value,
)
from app.reports.export.xlsx.analytics_overview import build_analytics_overview_sheet
from app.reports.export.xlsx.analytics_profiles import prepare_analytics_workbook_data
from app.reports.export.xlsx.analytics_schema import (
    ANALYTICS_DICTIONARY_SHEET,
    ANALYTICS_METADATA_SHEET,
    ANALYTICS_OVERVIEW_SHEET,
    ANALYTICS_WORKBOOK_FAMILY,
    ANALYTICS_WORKBOOK_STANDARD,
    AnalyticsColumn,
    AnalyticsExportContext,
    AnalyticsTableData,
    AnalyticsWorkbookData,
)
from app.reports.export.xlsx.schema import XLSX_REPORTING_STANDARD
from app.reports.export.xlsx.style import (
    KAORI_AMBER_FILL,
    KAORI_AMBER_TEXT,
    KAORI_BLUE_FILL,
    KAORI_BLUE_TEXT,
    KAORI_GREEN_FILL,
    KAORI_GREEN_TEXT,
    KAORI_RED_FILL,
    KAORI_RED_TEXT,
    KAORI_SLATE_FILL,
    KAORI_SLATE_TEXT,
    configure_sheet_page,
    set_footer,
    style_body,
    style_header,
    style_metadata_section,
    style_navigation_link,
    style_subtitle,
    style_title,
)

_TABLE_HEADER_ROW = 5
_TABLE_DATA_START_ROW = 6
_METADATA_HEADER_ROW = 5
_METADATA_DATA_START_ROW = 6
_DICTIONARY_HEADER_ROW = 5
_DICTIONARY_DATA_START_ROW = 6

_SECTION_LABELS = {'REPORT IDENTITY', 'REPORTING BOUNDARY', 'WORKBOOK CONTRACT', 'DATA HANDLING', 'SECURITY'}


def build_analytics_workbook(
    *,
    report_type: str,
    payload: Mapping[str, Any],
    context: AnalyticsExportContext,
    output_path: Path,
) -> None:
    data = prepare_analytics_workbook_data(report_type, payload, context)
    workbook = Workbook()
    overview = workbook.active
    overview.title = ANALYTICS_OVERVIEW_SHEET

    for table in data.tables:
        if table.sheet_name in {ANALYTICS_OVERVIEW_SHEET, ANALYTICS_METADATA_SHEET, ANALYTICS_DICTIONARY_SHEET}:
            raise ValueError(f'Analytics category sheet name {table.sheet_name!r} is reserved.')
        sheet = workbook.create_sheet(table.sheet_name)
        _build_evidence_sheet(sheet, table, data)

    metadata = workbook.create_sheet(ANALYTICS_METADATA_SHEET)
    _build_metadata_sheet(metadata, data)
    dictionary = workbook.create_sheet(ANALYTICS_DICTIONARY_SHEET)
    _build_dictionary_sheet(dictionary, data)
    build_analytics_overview_sheet(overview, data)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    validate_analytics_xlsx_artifact(output_path)


def validate_analytics_xlsx_artifact(output_path: Path) -> None:
    workbook = load_workbook(output_path, read_only=False, data_only=False)
    sheetnames = workbook.sheetnames
    if not sheetnames or sheetnames[0] != ANALYTICS_OVERVIEW_SHEET:
        raise ValueError('Analytics XLSX artifact must begin with Overview.')
    if sheetnames[-2:] != [ANALYTICS_METADATA_SHEET, ANALYTICS_DICTIONARY_SHEET]:
        raise ValueError('Analytics XLSX artifact must end with Metadata and Field Dictionary.')

    metadata = workbook[ANALYTICS_METADATA_SHEET]
    metadata_values = {
        metadata.cell(row=row, column=1).value: metadata.cell(row=row, column=2).value
        for row in range(1, metadata.max_row + 1)
        if metadata.cell(row=row, column=1).value
    }
    if metadata_values.get('Workbook Standard') != ANALYTICS_WORKBOOK_STANDARD:
        raise ValueError('Analytics XLSX artifact uses a stale or missing workbook standard.')
    if metadata_values.get('XLSX Standard') != XLSX_REPORTING_STANDARD:
        raise ValueError('Analytics XLSX artifact uses a stale or missing global XLSX standard.')
    if metadata_values.get('Report Family') != ANALYTICS_WORKBOOK_FAMILY:
        raise ValueError('Analytics XLSX artifact metadata must declare the analytics family.')

    overview = workbook[ANALYTICS_OVERVIEW_SHEET]
    if not str(overview['A1'].value or '').startswith('KAORI — '):
        raise ValueError('Analytics XLSX Overview title is missing.')
    if overview['A9'].value != 'Executive Readout':
        raise ValueError('Analytics XLSX Overview must contain the Executive Readout section.')
    navigation_targets = {
        str(cell.hyperlink.target)
        for row in overview.iter_rows(min_row=32, max_row=33)
        for cell in row
        if cell.hyperlink is not None
    }
    for sheet_name in sheetnames[1:]:
        if f"#'{sheet_name}'!A1" not in navigation_targets:
            raise ValueError(f'Analytics XLSX Overview is missing navigation to {sheet_name!r}.')

    for sheet_name in sheetnames[1:-2]:
        sheet = workbook[sheet_name]
        if sheet['A3'].hyperlink is None or sheet['A3'].hyperlink.target != "#'Overview'!A1":
            raise ValueError(f'Analytics XLSX sheet {sheet_name!r} must link back to Overview.')

    visual_header = overview['A13'].value
    if visual_header != 'Primary Visual':
        raise ValueError('Analytics XLSX Overview must reserve the Primary Visual surface.')
    support_header = overview['J13'].value
    if support_header != 'Supporting Data':
        raise ValueError('Analytics XLSX Overview must reserve the Supporting Data surface.')

    primary_visual = str(metadata_values.get('Primary Visual') or '').strip()
    if primary_visual and primary_visual != 'No chartable evidence' and not overview._charts:
        raise ValueError('Analytics XLSX Primary Visual is missing for chartable committed evidence.')

    if overview._charts:
        chart = overview._charts[0]
        chart_name = chart.__class__.__name__
        if chart_name == 'DoughnutChart':
            if chart.legend is None or chart.legend.position != 'b':
                raise ValueError('Analytics doughnut charts require a bottom legend.')
            labels = getattr(chart, 'dLbls', None)
            if labels is None or labels.showPercent is not True or labels.showVal is True or labels.showCatName is True:
                raise ValueError('Analytics doughnut charts require percentage-only direct labels.')
        if chart_name == 'BarChart':
            if getattr(chart, 'varyColors', None) is not False:
                raise ValueError('Analytics ranked/capacity bar charts must use controlled series colors.')
            series_count = len(chart.series)
            if series_count == 1 and chart.legend is not None:
                raise ValueError('Analytics single-series ranked bar charts must not use a redundant legend.')
            if series_count == 2:
                if chart.grouping != 'stacked' or chart.overlap != 100:
                    raise ValueError('Analytics capacity bars require a 100% stacked utilization/headroom presentation.')
                if chart.legend is None or chart.legend.position != 'b':
                    raise ValueError('Analytics capacity bars require a bottom legend.')
                labels = [str(getattr(series.tx, 'v', '') or '').strip() for series in chart.series]
                if labels != ['Utilization', 'Headroom']:
                    raise ValueError('Analytics capacity bars require Utilization and Headroom series labels.')
        if chart_name == 'LineChart':
            if chart.legend is None or chart.legend.position != 'b':
                raise ValueError('Analytics line charts require a bottom legend.')
            if any(not str(getattr(series.tx, 'v', '') or '').strip() for series in chart.series):
                raise ValueError('Analytics line charts require explicit series labels.')


def _build_evidence_sheet(sheet: Worksheet, table: AnalyticsTableData, data: AnalyticsWorkbookData) -> None:
    last_column = max(1, len(table.columns))
    last_letter = get_column_letter(last_column)
    sheet.merge_cells(f'A1:{last_letter}1')
    title = sheet['A1']
    title.value = f'KAORI — {table.title}'
    style_title(title)
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells(f'A2:{last_letter}2')
    subtitle = sheet['A2']
    subtitle.value = table.description
    style_subtitle(subtitle)
    sheet.row_dimensions[2].height = 28

    _write_back_to_overview(sheet)
    sheet.row_dimensions[3].height = 20
    sheet.row_dimensions[4].height = 8

    for column_index, column in enumerate(table.columns, start=1):
        cell = sheet.cell(row=_TABLE_HEADER_ROW, column=column_index, value=column.label)
        style_header(cell)
    sheet.row_dimensions[_TABLE_HEADER_ROW].height = 30

    for row_offset, raw_row in enumerate(table.rows):
        row_index = _TABLE_DATA_START_ROW + row_offset
        for column_index, column in enumerate(table.columns, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=_cell_value(raw_row.get(column.key), column))
            horizontal = 'right' if column.data_type in {'integer', 'number', 'percent'} else 'left'
            style_body(cell, alternate=row_offset % 2 == 1, horizontal=horizontal)
            if column.number_format:
                cell.number_format = column.number_format
            _apply_semantic_style(cell, column.label)
        sheet.row_dimensions[row_index].height = _row_height(table, raw_row)

    data_end = max(_TABLE_DATA_START_ROW, _TABLE_DATA_START_ROW + len(table.rows) - 1)
    table_name = _safe_table_name(f'{data.profile.report_type}_{table.sheet_name}')
    excel_table = Table(displayName=table_name, ref=f'A{_TABLE_HEADER_ROW}:{last_letter}{data_end}')
    excel_table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False, showFirstColumn=False, showLastColumn=False)
    sheet.add_table(excel_table)
    sheet.freeze_panes = f'A{_TABLE_DATA_START_ROW}'
    sheet.sheet_view.zoomScale = 85
    sheet.sheet_view.zoomScaleNormal = 85
    for index, column in enumerate(table.columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(column.width, 12), 48)
    configure_sheet_page(
        sheet,
        print_area=f'A1:{last_letter}{data_end}',
        print_title_rows=f'$1:${_TABLE_HEADER_ROW}',
        orientation='landscape' if last_column > 6 else 'portrait',
        fit_to_width=1 if last_column <= 12 else 2,
    )
    set_footer(sheet, str(data.context.report_id))


def _build_metadata_sheet(sheet: Worksheet, data: AnalyticsWorkbookData) -> None:
    sheet.merge_cells('A1:D1')
    title = sheet['A1']
    title.value = 'KAORI Analytics Workbook Metadata'
    style_title(title)
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells('A2:D2')
    subtitle = sheet['A2']
    subtitle.value = f'{data.profile.title} — publication boundary, workbook contract, and generation evidence.'
    style_subtitle(subtitle)
    sheet.row_dimensions[2].height = 28
    _write_back_to_overview(sheet)
    sheet.row_dimensions[3].height = 20
    sheet.row_dimensions[4].height = 8

    for column_index, label in enumerate(('Field', 'Value', 'Group', 'Description'), start=1):
        cell = sheet.cell(row=_METADATA_HEADER_ROW, column=column_index, value=label)
        style_header(cell)

    metadata = [
        ('REPORT IDENTITY', None, None, None),
        ('Report Reference', data.context.report_id, 'Identity', 'Canonical generated report reference.'),
        ('Report Name', data.context.report_name, 'Identity', 'Requested report name.'),
        ('Report Type', data.context.report_type, 'Identity', 'Canonical report catalog key.'),
        ('Report Family', ANALYTICS_WORKBOOK_FAMILY, 'Identity', 'Analytics workbook family.'),
        ('Classification', data.context.classification, 'Identity', 'Artifact classification.'),
        ('REPORTING BOUNDARY', None, None, None),
        ('Scope Type', data.context.scope_label, 'Boundary', 'Canonical infrastructure asset scope type.'),
        ('Scope', data.context.scope_value, 'Boundary', 'Selected infrastructure asset.'),
        ('Namespace', data.context.namespace_scope, 'Boundary', 'Selected namespace scope.'),
        ('Period', data.context.period_label, 'Boundary', 'Selected analytics period.'),
        ('Data Cut-off', data.context.data_cutoff, 'Boundary', 'Committed analytics evidence cut-off.'),
        ('Generated At', data.context.generated_at, 'Boundary', 'Artifact generation timestamp.'),
        ('Generated By', data.context.generated_by, 'Boundary', 'Stored generation identity snapshot.'),
        ('WORKBOOK CONTRACT', None, None, None),
        ('Workbook Standard', ANALYTICS_WORKBOOK_STANDARD, 'Contract', 'Analytics-specific XLSX workbook contract.'),
        ('XLSX Standard', XLSX_REPORTING_STANDARD, 'Contract', 'Global KAORI XLSX presentation standard.'),
        ('Tabular Standard', TABULAR_EXPORT_STANDARD, 'Contract', 'Shared typed-cell and formula-safety contract.'),
        ('Schema Version', TABULAR_SCHEMA_VERSION, 'Contract', 'Shared tabular schema version.'),
        ('Primary Format', 'XLSX', 'Contract', 'Native Excel analysis artifact; Analytics PDF remains the default publication.'),
        ('DATA HANDLING', None, None, None),
        ('Category Sheets', ', '.join(table.sheet_name for table in data.tables) or 'None', 'Data', 'Evidence sheets included only when committed category evidence exists.'),
        ('Primary Visual', data.primary_chart.visual_type if data.primary_chart else 'No chartable evidence', 'Data', 'Semantic visual selected by the Analytics report category.'),
        ('Recommendations', 'Backend Owned', 'Data', 'Excel never invents recommendations at export time.'),
        ('Forecast Calculation', 'Backend Owned', 'Data', 'Excel presents committed forecasts and does not calculate forecast models.'),
        ('SECURITY', None, None, None),
        ('Formula Injection Protection', 'Enabled', 'Security', 'Formula-like text values are escaped before workbook serialization.'),
    ]

    for row_index, row in enumerate(metadata, start=_METADATA_DATA_START_ROW):
        is_section = row[0] in _SECTION_LABELS
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=_metadata_value(value))
            if is_section:
                style_metadata_section(cell)
            else:
                style_body(cell, alternate=row_index % 2 == 0)
                if isinstance(value, datetime):
                    cell.number_format = 'yyyy-mm-dd hh:mm'
        sheet.row_dimensions[row_index].height = 22 if is_section else 27

    end_row = _METADATA_DATA_START_ROW + len(metadata) - 1
    table = Table(displayName='AnalyticsWorkbookMetadataTable', ref=f'A{_METADATA_HEADER_ROW}:D{end_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    sheet.add_table(table)
    sheet.freeze_panes = f'A{_METADATA_DATA_START_ROW}'
    sheet.sheet_view.zoomScale = 90
    for column, width in {'A': 30, 'B': 42, 'C': 18, 'D': 56}.items():
        sheet.column_dimensions[column].width = width
    configure_sheet_page(sheet, print_area=f'A1:D{end_row}', print_title_rows=f'$1:${_METADATA_HEADER_ROW}', fit_to_width=1)
    set_footer(sheet, str(data.context.report_id))


def _build_dictionary_sheet(sheet: Worksheet, data: AnalyticsWorkbookData) -> None:
    sheet.merge_cells('A1:I1')
    title = sheet['A1']
    title.value = 'KAORI Analytics Field Dictionary'
    style_title(title)
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells('A2:I2')
    subtitle = sheet['A2']
    subtitle.value = f'{data.profile.title} — category sheet fields, types, units, and evidence meaning.'
    style_subtitle(subtitle)
    sheet.row_dimensions[2].height = 28
    _write_back_to_overview(sheet)
    sheet.row_dimensions[3].height = 20
    sheet.row_dimensions[4].height = 8

    headers = ('Sheet', 'Field Key', 'Column', 'Definition', 'Data Type', 'Unit', 'Excel Number Format', 'Source', 'Missing Rule')
    for index, label in enumerate(headers, start=1):
        cell = sheet.cell(row=_DICTIONARY_HEADER_ROW, column=index, value=label)
        style_header(cell)

    rows: list[tuple[Any, ...]] = []
    for table in data.tables:
        for column in table.columns:
            rows.append(
                (
                    table.sheet_name,
                    column.key,
                    column.label,
                    column.definition,
                    column.data_type,
                    column.unit or 'None',
                    column.number_format or 'General',
                    'Canonical committed analytics payload',
                    'Not Available / blank when evidence is absent',
                )
            )
    if not rows:
        rows.append(('Overview', 'evidence_state', 'Evidence State', 'Explicit analytics evidence availability state.', 'text', 'None', 'General', 'Canonical committed analytics payload', 'Not Available'))

    for offset, row in enumerate(rows):
        row_index = _DICTIONARY_DATA_START_ROW + offset
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            style_body(cell, alternate=offset % 2 == 1)
        sheet.row_dimensions[row_index].height = 30

    end_row = _DICTIONARY_DATA_START_ROW + len(rows) - 1
    table = Table(displayName='AnalyticsFieldDictionaryTable', ref=f'A{_DICTIONARY_HEADER_ROW}:I{end_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    sheet.add_table(table)
    sheet.freeze_panes = f'A{_DICTIONARY_DATA_START_ROW}'
    sheet.sheet_view.zoomScale = 85
    for column, width in {'A': 20, 'B': 25, 'C': 28, 'D': 52, 'E': 16, 'F': 12, 'G': 20, 'H': 34, 'I': 34}.items():
        sheet.column_dimensions[column].width = width
    configure_sheet_page(sheet, print_area=f'A1:I{end_row}', print_title_rows=f'$1:${_DICTIONARY_HEADER_ROW}', fit_to_width=1)
    set_footer(sheet, str(data.context.report_id))


def _write_back_to_overview(sheet: Worksheet) -> None:
    cell = sheet['A3']
    cell.value = '← Overview'
    cell.hyperlink = "#'Overview'!A1"
    style_navigation_link(cell)


def _cell_value(value: Any, column: AnalyticsColumn) -> Any:
    if value in (None, ''):
        return None if column.data_type in {'integer', 'number', 'percent'} else 'Not Available'
    if column.data_type in {'integer', 'number', 'percent'}:
        number = _number(value)
        if number is not None:
            return int(number) if column.data_type == 'integer' and number.is_integer() else number
        return None
    if column.data_type in {'datetime', 'date'}:
        if isinstance(value, datetime):
            return _naive_datetime(value)
        if isinstance(value, date):
            return value
    if isinstance(value, (Mapping, list, tuple, set)):
        return protect_spreadsheet_formula(serialize_tabular_value(value))
    return protect_spreadsheet_formula(str(value))


def _metadata_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return _naive_datetime(value)
    if isinstance(value, date):
        return value
    if value is None:
        return None
    if isinstance(value, (Mapping, list, tuple, set)):
        return protect_spreadsheet_formula(serialize_tabular_value(value))
    return protect_spreadsheet_formula(str(value)) if isinstance(value, str) else value


def _number(value: Any) -> float | None:
    if isinstance(value, bool):
        return float(int(value))
    if isinstance(value, (int, float)):
        return float(value)
    if value is None:
        return None
    rendered = str(value).strip().replace(',', '').rstrip('%')
    try:
        return float(rendered)
    except ValueError:
        return None


def _naive_datetime(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value
    return value.astimezone(UTC).replace(tzinfo=None)


def _safe_table_name(value: str) -> str:
    cleaned = re.sub(r'[^A-Za-z0-9_]', '', ''.join(part.title() for part in value.replace('-', '_').split('_')))
    if not cleaned:
        cleaned = 'AnalyticsTable'
    if cleaned[0].isdigit():
        cleaned = f'T{cleaned}'
    return cleaned[:240]


def _row_height(table: AnalyticsTableData, row: Mapping[str, Any]) -> float:
    longest = max((len(str(row.get(column.key) or '')) for column in table.columns), default=0)
    lines = max(1, min(3, (longest // 56) + 1))
    return float(20 + ((lines - 1) * 10))


def _apply_semantic_style(cell: Any, label: str) -> None:
    if label.casefold() not in {'status', 'severity', 'priority', 'evidence state', 'signal quality'}:
        return
    value = str(cell.value or '').casefold()
    if any(token in value for token in ('critical', 'failed', 'error', 'unavailable', 'not ready')):
        cell.fill = PatternFill('solid', fgColor=KAORI_RED_FILL)
        cell.font = Font(name=cell.font.name, size=cell.font.sz, bold=True, color=KAORI_RED_TEXT)
    elif any(token in value for token in ('warning', 'attention', 'review', 'partial', 'pending')):
        cell.fill = PatternFill('solid', fgColor=KAORI_AMBER_FILL)
        cell.font = Font(name=cell.font.name, size=cell.font.sz, bold=True, color=KAORI_AMBER_TEXT)
    elif any(token in value for token in ('normal', 'healthy', 'ready', 'current', 'complete', 'available')):
        cell.fill = PatternFill('solid', fgColor=KAORI_GREEN_FILL)
        cell.font = Font(name=cell.font.name, size=cell.font.sz, bold=True, color=KAORI_GREEN_TEXT)
    elif 'info' in value:
        cell.fill = PatternFill('solid', fgColor=KAORI_BLUE_FILL)
        cell.font = Font(name=cell.font.name, size=cell.font.sz, bold=True, color=KAORI_BLUE_TEXT)
    else:
        cell.fill = PatternFill('solid', fgColor=KAORI_SLATE_FILL)
        cell.font = Font(name=cell.font.name, size=cell.font.sz, bold=True, color=KAORI_SLATE_TEXT)
