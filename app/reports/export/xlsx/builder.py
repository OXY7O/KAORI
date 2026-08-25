from __future__ import annotations

from collections.abc import Mapping
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, Final

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.export.inventory_common import (
    InventoryExportContext,
    InventoryExportDataset,
    build_inventory_export_dataset,
)
from app.reports.export.xlsx.overview import build_inventory_overview_sheet
from app.reports.export.tabular import (
    TABULAR_EXPORT_STANDARD,
    TABULAR_SCHEMA_VERSION,
    protect_spreadsheet_formula,
    serialize_tabular_value,
)
from app.reports.export.xlsx.standard import (
    INVENTORY_DATA_SHEET,
    INVENTORY_DICTIONARY_SHEET,
    INVENTORY_METADATA_SHEET,
    INVENTORY_OVERVIEW_SHEET,
    INVENTORY_WORKBOOK_FAMILY,
    INVENTORY_WORKBOOK_STANDARD,
    InventoryColumn,
    InventoryWorkbookProfile,
)
from app.reports.export.xlsx.style import (
    KAORI_AMBER_FILL,
    KAORI_AMBER_TEXT,
    KAORI_BLUE_FILL,
    KAORI_BLUE_TEXT,
    KAORI_GREEN_FILL,
    KAORI_GREEN_TEXT,
    KAORI_RED_FILL,
    KAORI_RED_TEXT,
    KAORI_SLATE_FILL,
    KAORI_SLATE_TEXT,
    configure_sheet_page,
    set_footer,
    style_body,
    style_context,
    style_empty,
    style_header,
    style_subtitle,
    style_title,
)

HEADER_ROW = 5
DATA_START_ROW = 6
_METADATA_HEADER_ROW: Final = 4
_METADATA_DATA_START_ROW: Final = 5
_DICTIONARY_HEADER_ROW: Final = 4
_DICTIONARY_DATA_START_ROW: Final = 5


_STATUS_GROUPS: Final = (
    (
        (
            'Ready',
            'Yes',
            'Current',
            'Valid',
            'In Use',
            'Fully Covered',
            'Passed',
            'Schedulable',
            'Discovered',
            'Active',
            'Bound',
            'Running',
            'Available',
        ),
        KAORI_GREEN_FILL,
        KAORI_GREEN_TEXT,
    ),
    (
        (
            'Review',
            'Warning',
            'Renewal Window',
            'Expiring Soon',
            'Unknown Usage',
            'Partial Coverage',
            'Stale',
            'Pending',
            'Terminating',
        ),
        KAORI_AMBER_FILL,
        KAORI_AMBER_TEXT,
    ),
    (
        (
            'Critical',
            'Critical Expiry',
            'Expired',
            'Invalid',
            'Failed',
            'Unavailable',
            'Not Ready',
            'No',
            'Lost',
            'Error',
        ),
        KAORI_RED_FILL,
        KAORI_RED_TEXT,
    ),
    (
        ('Unused', 'Information', 'Query Ready'),
        KAORI_BLUE_FILL,
        KAORI_BLUE_TEXT,
    ),
    (
        ('Unknown', 'Not Available', 'Not Collected', 'Not Reported', 'TLS Disabled'),
        KAORI_SLATE_FILL,
        KAORI_SLATE_TEXT,
    ),
)


def build_inventory_workbook(
    *,
    report_type: str,
    payload: Mapping[str, Any],
    context: InventoryExportContext,
    output_path: Path,
) -> None:
    dataset = build_inventory_export_dataset(
        report_type=report_type,
        payload=payload,
        context=context,
    )

    workbook = Workbook()
    summary_sheet = workbook.active
    if summary_sheet is None:
        summary_sheet = workbook.create_sheet(INVENTORY_OVERVIEW_SHEET)
    else:
        summary_sheet.title = INVENTORY_OVERVIEW_SHEET
    inventory_sheet = workbook.create_sheet(INVENTORY_DATA_SHEET)
    metadata_sheet = workbook.create_sheet(INVENTORY_METADATA_SHEET)
    dictionary_sheet = workbook.create_sheet(INVENTORY_DICTIONARY_SHEET)

    build_inventory_overview_sheet(summary_sheet, dataset)
    _build_inventory_sheet(inventory_sheet, dataset)
    _build_metadata_sheet(metadata_sheet, dataset, payload)
    _build_dictionary_sheet(dictionary_sheet, dataset.profile, context)

    workbook.properties.title = f'KAORI - {dataset.profile.title}'
    workbook.properties.subject = dataset.profile.description
    workbook.properties.creator = context.generated_by or 'KAORI'
    workbook.properties.keywords = f'KAORI,{INVENTORY_WORKBOOK_FAMILY},{report_type},xlsx,{TABULAR_EXPORT_STANDARD}'
    workbook.properties.created = _naive_datetime(context.generated_at)
    workbook.properties.modified = _naive_datetime(context.generated_at)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)




def _build_inventory_sheet(
    sheet: Worksheet,
    dataset: InventoryExportDataset,
) -> None:
    profile = dataset.profile
    max_column = len(profile.columns)
    last_column = get_column_letter(max_column)

    sheet.merge_cells(f'A1:{last_column}1')
    title = sheet['A1']
    title.value = f'KAORI — {profile.title}'
    style_title(title)
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells(f'A2:{last_column}2')
    subtitle = sheet['A2']
    subtitle.value = profile.description
    style_subtitle(subtitle)
    sheet.row_dimensions[2].height = 30

    sheet.merge_cells(f'A3:{last_column}3')
    context_cell = sheet['A3']
    context_cell.value = _inventory_context_line(dataset)
    style_context(context_cell)
    sheet.row_dimensions[3].height = 26
    sheet.row_dimensions[4].height = 8

    for index, column_spec in enumerate(profile.columns, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=column_spec.label)
        style_header(cell)
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(float(column_spec.width), 10.0),
            48.0,
        )
    sheet.row_dimensions[HEADER_ROW].height = 34

    if dataset.rows:
        for row_index, row in enumerate(dataset.rows, start=DATA_START_ROW):
            _write_inventory_row(sheet, row_index, row, profile)
    else:
        _write_empty_inventory_row(sheet, profile)

    final_row = max(DATA_START_ROW, DATA_START_ROW + dataset.row_count - 1)
    table = Table(
        displayName=f'{_safe_table_name(profile.report_type)}InventoryTable',
        ref=f'A{HEADER_ROW}:{last_column}{final_row}',
    )
    table.tableStyleInfo = TableStyleInfo(
        name='TableStyleMedium4',
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.auto_filter.ref = f'A{HEADER_ROW}:{last_column}{final_row}'

    frozen_column = get_column_letter(max(1, len(profile.identity_columns) + 1))
    sheet.freeze_panes = f'{frozen_column}{DATA_START_ROW}'

    _apply_status_conditional_formatting(
        sheet,
        profile,
        DATA_START_ROW,
        final_row,
    )
    configure_sheet_page(
        sheet,
        print_area=f'A1:{last_column}{final_row}',
        print_title_rows=f'$1:${HEADER_ROW}',
        fit_to_width=_inventory_print_width(max_column),
    )
    set_footer(sheet, str(dataset.context.report_id))


def _write_inventory_row(
    sheet: Worksheet,
    row_index: int,
    row: Mapping[str, Any],
    profile: InventoryWorkbookProfile,
) -> None:
    rendered_values: list[str] = []
    for column_index, column_spec in enumerate(profile.columns, start=1):
        value = _cell_value(row.get(column_spec.key), column_spec)
        cell = sheet.cell(row=row_index, column=column_index, value=value)
        style_body(cell, alternate=row_index % 2 == 0)
        if column_spec.number_format:
            cell.number_format = column_spec.number_format
        rendered_values.append('' if value is None else str(value))

    sheet.row_dimensions[row_index].height = _inventory_row_height(rendered_values)


def _write_empty_inventory_row(
    sheet: Worksheet,
    profile: InventoryWorkbookProfile,
) -> None:
    message = 'No committed inventory rows were found for the selected reporting boundary.'
    for column_index, column_spec in enumerate(profile.columns, start=1):
        value = message if column_index == 1 else None
        cell = sheet.cell(row=DATA_START_ROW, column=column_index, value=value)
        style_empty(cell, primary=column_index == 1)
        if column_spec.number_format:
            cell.number_format = column_spec.number_format
    sheet.row_dimensions[DATA_START_ROW].height = 28


def _build_metadata_sheet(
    sheet: Worksheet,
    dataset: InventoryExportDataset,
    payload: Mapping[str, Any],
) -> None:
    context = dataset.context
    profile = dataset.profile

    sheet.merge_cells('A1:D1')
    title = sheet['A1']
    title.value = 'KAORI Inventory Workbook Metadata'
    style_title(title)
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells('A2:D2')
    subtitle = sheet['A2']
    subtitle.value = 'Canonical source, scope, schema, security, and parity controls.'
    style_subtitle(subtitle)
    sheet.row_dimensions[2].height = 26
    sheet.row_dimensions[3].height = 8

    headers = ('Field', 'Value', 'Applies To', 'Notes')
    for index, label in enumerate(headers, start=1):
        cell = sheet.cell(row=_METADATA_HEADER_ROW, column=index, value=label)
        style_header(cell)
    sheet.row_dimensions[_METADATA_HEADER_ROW].height = 28

    metadata = [
        (
            'Workbook Standard',
            INVENTORY_WORKBOOK_STANDARD,
            'All sheets',
            'Canonical reusable KAORI Inventory Excel contract.',
        ),
        (
            'Tabular Standard',
            TABULAR_EXPORT_STANDARD,
            'Workbook',
            'Shared CSV/XLSX source, identity, and metadata contract.',
        ),
        ('Schema Version', TABULAR_SCHEMA_VERSION, 'Workbook', 'Public tabular schema version.'),
        (
            'Report Reference',
            str(context.report_id),
            'Workbook',
            'Stable KAORI document reference; internal database keys are not exposed.',
        ),
        ('Report Name', context.report_name, 'Workbook', 'User-visible report request name.'),
        ('Report Type', context.report_type, 'Workbook', profile.title),
        (
            'Report Family',
            INVENTORY_WORKBOOK_FAMILY,
            'Workbook',
            'Inventory is table-first and excludes analytical recommendations.',
        ),
        (
            'Primary Format',
            'XLSX',
            'Workbook',
            'Native Excel workbook with Overview, canonical Inventory, Metadata, and Field Dictionary sheets.',
        ),
        ('CSV Parity', 'Enabled', 'Inventory', 'CSV and XLSX use the same InventoryExportDataset and profile order.'),
        ('Classification', context.classification, 'Workbook', 'Artifact handling classification.'),
        ('Generated At', context.generated_at, 'Workbook', 'Artifact generation timestamp.'),
        ('Generated By', context.generated_by, 'Workbook', 'Committed user display-name or username snapshot.'),
        (
            f'Selected {context.scope_label}',
            context.cluster_scope,
            'Inventory',
            'Canonical infrastructure asset boundary for this report.',
        ),
        (
            'Selected Namespace',
            context.namespace_scope,
            'Inventory',
            'All Namespaces remains inside the selected cluster.',
        ),
        (
            'Reporting Period',
            context.period_label,
            'Inventory',
            'Inventory reports use the committed snapshot boundary.',
        ),
        (
            'Data Cut-off',
            context.data_cutoff or payload.get('data_cutoff') or payload.get('collected_at'),
            'Inventory',
            'Committed evidence cut-off.',
        ),
        ('Source', context.source_label, 'Inventory', 'No live Kubernetes API call occurs during workbook rendering.'),
        ('Total Rows', dataset.row_count, 'Inventory', 'Number of committed inventory rows exported.'),
        (
            'Identity Columns',
            ', '.join(profile.identity_columns) or 'None',
            'Inventory',
            'Leading columns used for stable record identity and frozen-column behavior.',
        ),
        (
            'Missing Text Rule',
            'Profile-defined; normally Not Available',
            'Inventory',
            'Missing text evidence remains explicit.',
        ),
        ('Missing Numeric Rule', 'Blank', 'Inventory', 'Numeric columns remain numeric and sortable.'),
        ('Horizontal Alignment', 'Left', 'Inventory', 'Canonical KAORI table alignment.'),
        ('Vertical Alignment', 'Middle', 'Inventory', 'Text is vertically centered in every table row.'),
        (
            'Private Key Export',
            'Excluded',
            'TLS Inventory',
            'tls.key, Secret values, and recoverable key material are never exported.',
        ),
    ]

    for row_index, metadata_row in enumerate(metadata, start=_METADATA_DATA_START_ROW):
        for column_index, value in enumerate(metadata_row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=_metadata_value(value))
            style_body(cell, alternate=row_index % 2 == 0)
            if isinstance(value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm'
        sheet.row_dimensions[row_index].height = 28

    end_row = _METADATA_HEADER_ROW + len(metadata)
    table = Table(
        displayName='InventoryWorkbookMetadataTable',
        ref=f'A{_METADATA_HEADER_ROW}:D{end_row}',
    )
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=False)
    sheet.add_table(table)
    sheet.freeze_panes = f'A{_METADATA_DATA_START_ROW}'
    for column, width in {'A': 25, 'B': 42, 'C': 22, 'D': 58}.items():
        sheet.column_dimensions[column].width = width
    configure_sheet_page(
        sheet,
        print_area=f'A1:D{end_row}',
        print_title_rows=f'$1:${_METADATA_HEADER_ROW}',
        fit_to_width=1,
    )
    set_footer(sheet, str(context.report_id))


def _build_dictionary_sheet(
    sheet: Worksheet,
    profile: InventoryWorkbookProfile,
    context: InventoryExportContext,
) -> None:
    sheet.merge_cells('A1:I1')
    title = sheet['A1']
    title.value = 'KAORI Inventory Field Dictionary'
    style_title(title)
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells('A2:I2')
    subtitle = sheet['A2']
    subtitle.value = f'{profile.title} — canonical field order, meaning, source, and formatting.'
    style_subtitle(subtitle)
    sheet.row_dimensions[2].height = 26
    sheet.row_dimensions[3].height = 8

    headers = (
        'Field Key',
        'Column',
        'Definition',
        'Data Type',
        'Source / Rule',
        'Missing Value Rule',
        'Semantic Role',
        'Unit',
        'Excel Number Format',
    )
    for index, label in enumerate(headers, start=1):
        cell = sheet.cell(row=_DICTIONARY_HEADER_ROW, column=index, value=label)
        style_header(cell)
    sheet.row_dimensions[_DICTIONARY_HEADER_ROW].height = 30

    for row_index, column_spec in enumerate(profile.columns, start=_DICTIONARY_DATA_START_ROW):
        values = (
            column_spec.key,
            column_spec.label,
            column_spec.definition,
            column_spec.data_type,
            column_spec.source,
            column_spec.missing_value_rule,
            column_spec.semantic_type or 'None',
            column_spec.unit or 'None',
            column_spec.number_format or 'General',
        )
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            style_body(cell, alternate=row_index % 2 == 0)
        sheet.row_dimensions[row_index].height = 32

    end_row = _DICTIONARY_HEADER_ROW + len(profile.columns)
    table = Table(
        displayName=f'{_safe_table_name(profile.report_type)}FieldDictionaryTable',
        ref=f'A{_DICTIONARY_HEADER_ROW}:I{end_row}',
    )
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=False)
    sheet.add_table(table)
    sheet.freeze_panes = f'A{_DICTIONARY_DATA_START_ROW}'
    for column, width in {
        'A': 25,
        'B': 28,
        'C': 58,
        'D': 16,
        'E': 36,
        'F': 22,
        'G': 20,
        'H': 14,
        'I': 20,
    }.items():
        sheet.column_dimensions[column].width = width
    configure_sheet_page(
        sheet,
        print_area=f'A1:I{end_row}',
        print_title_rows=f'$1:${_DICTIONARY_HEADER_ROW}',
        fit_to_width=1,
    )
    set_footer(sheet, str(context.report_id))


def _inventory_context_line(dataset: InventoryExportDataset) -> str:
    context = dataset.context
    cutoff = serialize_tabular_value(context.data_cutoff) or 'Not Available'
    return (
        f'Reference: {context.report_id}  |  {context.scope_label}: {context.cluster_scope}  |  '
        f'Namespace: {context.namespace_scope}  |  Data Cut-off: {cutoff}  |  '
        f'Rows: {dataset.row_count}'
    )


def _inventory_print_width(column_count: int) -> int:
    if column_count <= 12:
        return 1
    if column_count <= 24:
        return 2
    return 3


def _inventory_row_height(values: list[str]) -> float:
    longest = max((len(value) for value in values), default=0)
    estimated_lines = max(1, min(4, (longest // 48) + 1))
    return float(min(60, 18 + ((estimated_lines - 1) * 12)))


def _cell_value(value: Any, column_spec: InventoryColumn) -> Any:
    if value is None or value == '':
        if column_spec.data_type in {'integer', 'number', 'decimal'}:
            return None
        return column_spec.missing_value_rule or 'Not Available'

    if column_spec.data_type == 'boolean':
        if isinstance(value, bool):
            return 'Yes' if value else 'No'
        return _safe_text_value(value)

    if column_spec.data_type in {'integer', 'number', 'decimal'}:
        return _numeric_value(value)

    if column_spec.data_type in {'datetime', 'date'}:
        if isinstance(value, datetime):
            return _naive_datetime(value)
        if isinstance(value, date):
            return value
        return _safe_text_value(value)

    if isinstance(value, (Mapping, list, tuple, set)):
        return protect_spreadsheet_formula(serialize_tabular_value(value))
    return _safe_text_value(value)


def _safe_text_value(value: Any) -> str:
    return protect_spreadsheet_formula(str(value))


def _numeric_value(value: Any) -> int | float | None:
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return value
    rendered = str(value).strip()
    if not rendered or rendered.lower() in {'not available', 'unknown', '-'}:
        return None
    try:
        number = float(rendered)
    except ValueError:
        return None
    return int(number) if number.is_integer() else number


def _metadata_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return _naive_datetime(value)
    if isinstance(value, date):
        return value
    if value is None or value == '':
        return 'Not Available'
    if isinstance(value, (Mapping, list, tuple, set)):
        return protect_spreadsheet_formula(serialize_tabular_value(value))
    return _safe_text_value(value) if isinstance(value, str) else value


def _naive_datetime(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value
    return value.astimezone(UTC).replace(tzinfo=None)


def _safe_table_name(report_type: str) -> str:
    return ''.join(part.title() for part in report_type.replace('-', '_').split('_'))


def _apply_status_conditional_formatting(
    sheet: Worksheet,
    profile: InventoryWorkbookProfile,
    start_row: int,
    end_row: int,
) -> None:
    for column_index, column_spec in enumerate(profile.columns, start=1):
        semantic_type = column_spec.semantic_type
        if semantic_type is None:
            continue

        letter = get_column_letter(column_index)
        range_ref = f'{letter}{start_row}:{letter}{max(start_row, end_row)}'
        first_cell = f'{letter}{start_row}'

        if semantic_type == 'negative_condition':
            _add_value_group_rule(
                sheet,
                range_ref,
                first_cell,
                ('True', 'Yes', 'Active'),
                KAORI_RED_FILL,
                KAORI_RED_TEXT,
            )
            _add_value_group_rule(
                sheet,
                range_ref,
                first_cell,
                ('False', 'No', 'Inactive'),
                KAORI_GREEN_FILL,
                KAORI_GREEN_TEXT,
            )
            continue

        for values, fill, font in _STATUS_GROUPS:
            _add_value_group_rule(sheet, range_ref, first_cell, values, fill, font)

    days_index = next(
        (index for index, column in enumerate(profile.columns, start=1) if column.key == 'days_remaining'),
        None,
    )
    if days_index is None:
        return

    letter = get_column_letter(days_index)
    range_ref = f'{letter}{start_row}:{letter}{max(start_row, end_row)}'
    sheet.conditional_formatting.add(
        range_ref,
        CellIsRule(
            operator='lessThan',
            formula=['0'],
            fill=PatternFill('solid', fgColor=KAORI_RED_FILL),
            font=Font(color=KAORI_RED_TEXT, bold=True),
        ),
    )
    sheet.conditional_formatting.add(
        range_ref,
        CellIsRule(
            operator='between',
            formula=['0', '7'],
            fill=PatternFill('solid', fgColor=KAORI_RED_FILL),
            font=Font(color=KAORI_RED_TEXT, bold=True),
        ),
    )
    sheet.conditional_formatting.add(
        range_ref,
        CellIsRule(
            operator='between',
            formula=['8', '30'],
            fill=PatternFill('solid', fgColor=KAORI_AMBER_FILL),
            font=Font(color=KAORI_AMBER_TEXT, bold=True),
        ),
    )


def _add_value_group_rule(
    sheet: Worksheet,
    range_ref: str,
    first_cell: str,
    values: tuple[str, ...],
    fill: str,
    font: str,
) -> None:
    expression = ','.join(f'"{value}"' for value in values)
    formula = f'ISNUMBER(MATCH({first_cell},{{{expression}}},0))'
    sheet.conditional_formatting.add(
        range_ref,
        FormulaRule(
            formula=[formula],
            fill=PatternFill('solid', fgColor=fill),
            font=Font(color=font, bold=True),
        ),
    )
