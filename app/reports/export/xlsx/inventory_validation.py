from __future__ import annotations

from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.chart import BarChart, DoughnutChart
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.export.xlsx.standard import (
    INVENTORY_DATA_SHEET,
    INVENTORY_DICTIONARY_SHEET,
    INVENTORY_METADATA_SHEET,
    INVENTORY_OVERVIEW_SHEET,
    INVENTORY_WORKBOOK_STANDARD,
    XLSX_REPORTING_STANDARD,
)

_METADATA_DATA_START_ROW = 5


def validate_inventory_xlsx_artifact(output_path: Path) -> None:
    """Reject serialized XLSX artifacts that drift from the global KAORI contract."""

    workbook = load_workbook(output_path, data_only=False, read_only=False)
    try:
        _validate_sheet_order(workbook.sheetnames)
        overview = cast(Worksheet, workbook[INVENTORY_OVERVIEW_SHEET])
        _validate_overview_structure(overview)
        _validate_overview_visual(overview)
        _validate_navigation(workbook, overview)
        _validate_metadata(cast(Worksheet, workbook[INVENTORY_METADATA_SHEET]))
    finally:
        workbook.close()


def _validate_sheet_order(sheetnames: list[str]) -> None:
    expected_sheets = [
        INVENTORY_OVERVIEW_SHEET,
        INVENTORY_DATA_SHEET,
        INVENTORY_METADATA_SHEET,
        INVENTORY_DICTIONARY_SHEET,
    ]
    if sheetnames != expected_sheets:
        raise ValueError(
            'KAORI XLSX contract violation: Overview, Inventory, Metadata, and '
            'Field Dictionary must be the canonical sheet order.'
        )


def _validate_overview_structure(overview: Worksheet) -> None:
    if overview['A5'].value != 'Executive Snapshot':
        raise ValueError('KAORI XLSX contract violation: Executive Snapshot is missing.')
    if overview['A9'].value != 'Executive Readout':
        raise ValueError('KAORI XLSX contract violation: Executive Readout is missing.')
    if overview['A13'].value != 'Primary Visual' or overview['I13'].value != 'Supporting Data':
        raise ValueError('KAORI XLSX contract violation: Primary Visual or Supporting Data section is missing.')


def _validate_overview_visual(overview: Worksheet) -> None:
    has_no_chart_message = str(overview['A15'].value or '').strip() == (
        'No committed chartable evidence is available for this workbook scope.'
    )
    overview_charts = cast(Any, overview)._charts
    if not overview_charts:
        if not has_no_chart_message:
            raise ValueError(
                'KAORI XLSX contract violation: Overview must contain an evidence-backed chart '
                'when chartable report evidence is available.'
            )
        return

    chart = overview_charts[0]
    if isinstance(chart, BarChart):
        _validate_bar_chart(chart)
    elif isinstance(chart, DoughnutChart):
        _validate_doughnut_chart(chart)
    else:
        raise ValueError('KAORI XLSX contract violation: unsupported Overview chart type.')

    _validate_supporting_data(overview)


def _validate_bar_chart(chart: BarChart) -> None:
    if getattr(chart, 'varyColors', None) is not False:
        raise ValueError('KAORI XLSX contract violation: ranked bar charts must set varyColors=False.')
    if chart.legend is not None:
        raise ValueError('KAORI XLSX contract violation: ranked bar charts must not use a redundant legend.')
    if chart.dLbls is None or chart.dLbls.showVal is not True:
        raise ValueError('KAORI XLSX contract violation: ranked bar charts must use direct value labels.')


def _validate_doughnut_chart(chart: DoughnutChart) -> None:
    if chart.legend is None or chart.legend.position != 'b':
        raise ValueError('KAORI XLSX contract violation: doughnut legend must be present at the bottom.')
    if (
        chart.dLbls is None
        or chart.dLbls.showPercent is not True
        or chart.dLbls.showLegendKey is not False
        or chart.dLbls.showVal is not False
        or chart.dLbls.showCatName is not False
        or chart.dLbls.showSerName is not False
        or chart.dLbls.showLeaderLines is not True
        or chart.dLbls.dLblPos != 'bestFit'
    ):
        raise ValueError(
            'KAORI XLSX contract violation: doughnut charts must use clean percent-only labels with leader lines.'
        )


def _validate_supporting_data(overview: Worksheet) -> None:
    if (
        overview['I15'].value != 'Rank'
        or overview['J15'].value in (None, '')
        or overview['L15'].value != 'Share of Total'
    ):
        raise ValueError('KAORI XLSX contract violation: chartable Overview must include canonical supporting data.')


def _validate_navigation(workbook: Any, overview: Worksheet) -> None:
    expected_navigation = {
        'A32': (INVENTORY_DATA_SHEET, f"#'{INVENTORY_DATA_SHEET}'!A1"),
        'E32': (INVENTORY_METADATA_SHEET, f"#'{INVENTORY_METADATA_SHEET}'!A1"),
        'I32': (INVENTORY_DICTIONARY_SHEET, f"#'{INVENTORY_DICTIONARY_SHEET}'!A1"),
    }
    for coordinate, (label, target) in expected_navigation.items():
        cell = overview[coordinate]
        actual_target = cell.hyperlink.target if cell.hyperlink is not None else None
        if cell.value != f'Open {label} →' or actual_target != target:
            raise ValueError(f'KAORI XLSX contract violation: Overview navigation for {label} is missing or invalid.')

    for sheet_name, coordinate in (
        (INVENTORY_DATA_SHEET, 'A4'),
        (INVENTORY_METADATA_SHEET, 'A3'),
        (INVENTORY_DICTIONARY_SHEET, 'A3'),
    ):
        cell = workbook[sheet_name][coordinate]
        actual_target = cell.hyperlink.target if cell.hyperlink is not None else None
        if cell.value != '← Overview' or actual_target != "#'Overview'!A1":
            raise ValueError(f'KAORI XLSX contract violation: {sheet_name} must link back to Overview.')


def _validate_metadata(metadata: Worksheet) -> None:
    metadata_values = {
        metadata.cell(row=row, column=1).value: metadata.cell(row=row, column=2).value
        for row in range(_METADATA_DATA_START_ROW, metadata.max_row + 1)
    }
    if metadata_values.get('Workbook Standard') != INVENTORY_WORKBOOK_STANDARD:
        raise ValueError('KAORI XLSX contract violation: Workbook Standard is stale or missing.')
    if metadata_values.get('XLSX Standard') != XLSX_REPORTING_STANDARD:
        raise ValueError('KAORI XLSX contract violation: XLSX Standard is stale or missing.')
