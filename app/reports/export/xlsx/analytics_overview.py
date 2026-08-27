from __future__ import annotations

from typing import Any, cast

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.export.xlsx.analytics_schema import AnalyticsChartData, AnalyticsWorkbookData
from app.reports.export.xlsx.style import (
    KAORI_ACCENT,
    KAORI_AMBER_TEXT,
    KAORI_BLUE_TEXT,
    KAORI_GREEN_TEXT,
    KAORI_PRIMARY,
    KAORI_PRIMARY_DARK,
    KAORI_PRIMARY_SOFT,
    KAORI_RED_TEXT,
    KAORI_SLATE_TEXT,
    configure_sheet_page,
    set_footer,
    style_body,
    style_context,
    style_header,
    style_navigation_link,
    style_readout,
    style_section_header,
    style_subtitle,
    style_summary_label,
    style_summary_value,
    style_title,
)

_OVERVIEW_COLUMNS = 14  # A:N
_SNAPSHOT_HEADER_ROW = 5
_KPI_LABEL_ROW = 6
_KPI_VALUE_ROW = 7
_READOUT_HEADER_ROW = 9
_READOUT_VALUE_ROW = 10
_VISUAL_HEADER_ROW = 13
_SUPPORTING_HEADER_ROW = 13
_SUPPORTING_START_COLUMN = 10  # J
_CHART_CATEGORY_COLUMN = 15  # O, hidden helper used by long-label ranked charts
_SUPPORTING_HEADER_DATA_ROW = 15
_SUPPORTING_DATA_ROW = 16
_NAV_HEADER_ROW = 31
_NAV_ROW = 32
_LAST_ROW = 34

_LINE_COLORS = (KAORI_PRIMARY, KAORI_BLUE_TEXT, KAORI_GREEN_TEXT, KAORI_AMBER_TEXT)
_DOUGHNUT_COLORS = (KAORI_RED_TEXT, KAORI_AMBER_TEXT, KAORI_BLUE_TEXT, KAORI_GREEN_TEXT, KAORI_ACCENT, KAORI_SLATE_TEXT)


def build_analytics_overview_sheet(sheet: Worksheet, data: AnalyticsWorkbookData) -> None:
    _write_header(sheet, data)
    _write_metrics(sheet, data)
    _write_readout(sheet, data)
    _write_visual_and_supporting_data(sheet, data)
    _write_navigation(sheet)

    for column, width in {
        'A': 12,
        'B': 12,
        'C': 12,
        'D': 12,
        'E': 12,
        'F': 12,
        'G': 12,
        'H': 12,
        'I': 4,
        'J': 8,
        'K': 28,
        'L': 16,
        'M': 16,
        'N': 16,
    }.items():
        sheet.column_dimensions[column].width = width
    if data.primary_chart is not None and data.primary_chart.visual_type == 'comparison_bar':
        for column, width in {'J': 20, 'K': 14, 'L': 14, 'M': 17, 'N': 20}.items():
            sheet.column_dimensions[column].width = width
    elif data.primary_chart is not None and data.primary_chart.visual_type == 'capacity_bar':
        for column, width in {'J': 20, 'K': 16, 'L': 16, 'M': 8, 'N': 8}.items():
            sheet.column_dimensions[column].width = width
    elif data.primary_chart is not None and data.primary_chart.visual_type == 'line':
        for column, width in {'J': 22, 'K': 14, 'L': 14, 'M': 14, 'N': 14}.items():
            sheet.column_dimensions[column].width = width
    elif data.primary_chart is not None and data.primary_chart.visual_type == 'ranked_bar':
        for column, width in {'J': 8, 'K': 38, 'L': 20, 'M': 16, 'N': 4}.items():
            sheet.column_dimensions[column].width = width
        sheet.column_dimensions['O'].hidden = True
    elif data.primary_chart is not None and data.primary_chart.visual_type == 'composition_bar':
        for column, width in {'J': 8, 'K': 30, 'L': 16, 'M': 16, 'N': 4}.items():
            sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A5'
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_view.zoomScaleNormal = 90
    configure_sheet_page(
        sheet,
        print_area=f'A1:N{_LAST_ROW}',
        print_title_rows='$1:$3',
        orientation='landscape',
        fit_to_width=1,
    )
    set_footer(sheet, str(data.context.report_id))


def _write_header(sheet: Worksheet, data: AnalyticsWorkbookData) -> None:
    sheet.merge_cells('A1:N1')
    title = sheet['A1']
    title.value = _overview_title(data)
    style_title(title)
    sheet.row_dimensions[1].height = 32

    sheet.merge_cells('A2:N2')
    subtitle = sheet['A2']
    subtitle.value = data.profile.description
    style_subtitle(subtitle)
    sheet.row_dimensions[2].height = 30

    sheet.merge_cells('A3:N3')
    context = sheet['A3']
    context.value = _context_line(data)
    style_context(context)
    sheet.row_dimensions[3].height = 24
    sheet.row_dimensions[4].height = 8


def _write_metrics(sheet: Worksheet, data: AnalyticsWorkbookData) -> None:
    sheet.merge_cells(start_row=_SNAPSHOT_HEADER_ROW, start_column=1, end_row=_SNAPSHOT_HEADER_ROW, end_column=14)
    header = sheet.cell(row=_SNAPSHOT_HEADER_ROW, column=1, value='Executive Snapshot')
    style_section_header(header)
    sheet.row_dimensions[_SNAPSHOT_HEADER_ROW].height = 22

    spans = ((1, 3), (4, 6), (7, 9), (10, 12))
    for metric, (start, end) in zip(data.metrics[:4], spans, strict=True):
        _merged_card(sheet, _KPI_LABEL_ROW, start, end, metric.label, style_summary_label)
        _merged_card(sheet, _KPI_VALUE_ROW, start, end, metric.value, style_summary_value)
        cell = sheet.cell(row=_KPI_VALUE_ROW, column=start)
        if metric.number_format:
            cell.number_format = metric.number_format
    sheet.row_dimensions[_KPI_LABEL_ROW].height = 22
    sheet.row_dimensions[_KPI_VALUE_ROW].height = 34
    sheet.row_dimensions[8].height = 8


def _write_readout(sheet: Worksheet, data: AnalyticsWorkbookData) -> None:
    sheet.merge_cells(start_row=_READOUT_HEADER_ROW, start_column=1, end_row=_READOUT_HEADER_ROW, end_column=14)
    header = sheet.cell(row=_READOUT_HEADER_ROW, column=1, value='Executive Readout')
    style_section_header(header)
    sheet.row_dimensions[_READOUT_HEADER_ROW].height = 22

    sheet.merge_cells(start_row=_READOUT_VALUE_ROW, start_column=1, end_row=_READOUT_VALUE_ROW + 1, end_column=14)
    cell = sheet.cell(row=_READOUT_VALUE_ROW, column=1, value=data.readout)
    style_readout(cell)
    sheet.row_dimensions[_READOUT_VALUE_ROW].height = 24
    sheet.row_dimensions[_READOUT_VALUE_ROW + 1].height = 24
    sheet.row_dimensions[12].height = 8


def _write_visual_and_supporting_data(sheet: Worksheet, data: AnalyticsWorkbookData) -> None:
    sheet.merge_cells(start_row=_VISUAL_HEADER_ROW, start_column=1, end_row=_VISUAL_HEADER_ROW, end_column=8)
    visual_header = sheet.cell(row=_VISUAL_HEADER_ROW, column=1, value='Primary Visual')
    style_section_header(visual_header)

    sheet.merge_cells(start_row=_SUPPORTING_HEADER_ROW, start_column=10, end_row=_SUPPORTING_HEADER_ROW, end_column=14)
    supporting_header = sheet.cell(row=_SUPPORTING_HEADER_ROW, column=10, value='Supporting Data')
    style_section_header(supporting_header)
    sheet.row_dimensions[_VISUAL_HEADER_ROW].height = 22
    sheet.row_dimensions[14].height = 10

    chart = data.primary_chart
    if chart is None or not chart.rows:
        sheet.merge_cells(start_row=15, start_column=1, end_row=18, end_column=8)
        message = sheet.cell(
            row=15,
            column=1,
            value='No committed chartable analytics evidence is available for this reporting boundary.',
        )
        style_readout(message)
        sheet.merge_cells(start_row=15, start_column=10, end_row=18, end_column=14)
        support = sheet.cell(
            row=15, column=10, value='Use the category evidence sheets for available committed records.'
        )
        style_body(support)
        return

    _write_supporting_data(sheet, chart)
    _add_chart(sheet, data, chart)


def _write_supporting_data(sheet: Worksheet, chart: AnalyticsChartData) -> None:
    if chart.visual_type == 'line':
        category_label = 'Observed At' if chart.category_key == 'timestamp' else 'Observation'
        headers = [category_label, *[series.label for series in chart.series[:4]]]
        for index, label in enumerate(headers, start=_SUPPORTING_START_COLUMN):
            cell = sheet.cell(row=_SUPPORTING_HEADER_DATA_ROW, column=index, value=label)
            style_header(cell)
        rows = list(chart.rows[-6:])
        for offset, row in enumerate(rows, start=0):
            row_index = _SUPPORTING_DATA_ROW + offset
            category = row.get(chart.category_key)
            line_values: list[Any] = [category]
            line_values.extend(row.get(series.key) for series in chart.series[:4])
            for column_index, value in enumerate(line_values, start=_SUPPORTING_START_COLUMN):
                cell = sheet.cell(row=row_index, column=column_index, value=value)
                style_body(
                    cell,
                    alternate=offset % 2 == 1,
                    horizontal='right' if column_index > _SUPPORTING_START_COLUMN else 'left',
                )
                series_index = column_index - _SUPPORTING_START_COLUMN - 1
                if series_index >= 0 and series_index < len(chart.series):
                    series = chart.series[series_index]
                    cell.number_format = _series_number_format(series.number_format, series.unit)
            sheet.row_dimensions[row_index].height = 25
        return

    if chart.visual_type == 'comparison_bar':
        headers = ('Resource', 'Current', 'Projected', 'Change (pp)', 'Trend')
        for index, label in enumerate(headers, start=_SUPPORTING_START_COLUMN):
            cell = sheet.cell(row=_SUPPORTING_HEADER_DATA_ROW, column=index, value=label)
            style_header(cell)
        for offset, row in enumerate(chart.rows[:8]):
            row_index = _SUPPORTING_DATA_ROW + offset
            comparison_values = (
                row.get(chart.category_key),
                _number(row.get('current_percent')),
                _number(row.get('projected_percent')),
                _number(row.get('change_pp')),
                row.get('trend'),
            )
            for column_index, item in enumerate(comparison_values, start=_SUPPORTING_START_COLUMN):
                cell = sheet.cell(row=row_index, column=column_index, value=item)
                horizontal = 'center' if column_index == 13 else ('right' if column_index in {11, 12} else 'left')
                style_body(
                    cell,
                    alternate=offset % 2 == 1,
                    horizontal=horizontal,
                )
            sheet.cell(row=row_index, column=11).number_format = '0.0"%"'
            sheet.cell(row=row_index, column=12).number_format = '0.0"%"'
            sheet.cell(row=row_index, column=13).number_format = '+0.0;-0.0;0.0'
            sheet.row_dimensions[row_index].height = 25
        return

    if chart.visual_type == 'capacity_bar':
        headers = ('Resource', 'Utilization', 'Headroom')
        for index, label in enumerate(headers, start=_SUPPORTING_START_COLUMN):
            cell = sheet.cell(row=_SUPPORTING_HEADER_DATA_ROW, column=index, value=label)
            style_header(cell)
        for offset, row in enumerate(chart.rows[:8]):
            row_index = _SUPPORTING_DATA_ROW + offset
            value = _number(row.get('value'))
            headroom = max(0.0, 100.0 - value) if value is not None else None
            capacity_values = (row.get(chart.category_key), value, headroom)
            for column_index, item in enumerate(capacity_values, start=_SUPPORTING_START_COLUMN):
                cell = sheet.cell(row=row_index, column=column_index, value=item)
                style_body(
                    cell,
                    alternate=offset % 2 == 1,
                    horizontal='right' if column_index > _SUPPORTING_START_COLUMN else 'left',
                )
            sheet.cell(row=row_index, column=11).number_format = '0.0"%"'
            sheet.cell(row=row_index, column=12).number_format = '0.0"%"'
            sheet.row_dimensions[row_index].height = 25
        return

    share_label = (
        'Share of Total' if chart.visual_type == 'composition_bar' or chart.total_value else 'Share of Displayed'
    )
    headers = ('Rank', 'Category', 'Value', share_label)
    for index, label in enumerate(headers, start=_SUPPORTING_START_COLUMN):
        cell = sheet.cell(row=_SUPPORTING_HEADER_DATA_ROW, column=index, value=label)
        style_header(cell)
    displayed_total = sum(_number(row.get('value')) for row in chart.rows)
    total = float(chart.total_value or 0) or displayed_total
    for rank, row in enumerate(chart.rows[:8], start=1):
        row_index = _SUPPORTING_DATA_ROW + rank - 1
        value = _number(row.get('value'))
        share = value / total * 100.0 if total > 0 else 0.0
        ranked_values = (rank, row.get(chart.category_key), value, share)
        for column_index, item in enumerate(ranked_values, start=_SUPPORTING_START_COLUMN):
            cell = sheet.cell(row=row_index, column=column_index, value=item)
            style_body(cell, alternate=rank % 2 == 0, horizontal='right' if column_index in {10, 12, 13} else 'left')
        if chart.series:
            sheet.cell(row=row_index, column=12).number_format = _series_number_format(
                chart.series[0].number_format,
                chart.series[0].unit,
            )
        sheet.cell(row=row_index, column=13).number_format = '0.0"%"'
        category = str(row.get(chart.category_key) or '')
        if chart.visual_type == 'ranked_bar':
            sheet.cell(row=row_index, column=_CHART_CATEGORY_COLUMN, value=_short_chart_label(category))
        sheet.row_dimensions[row_index].height = 34 if len(category) > 36 else 25


def _add_chart(sheet: Worksheet, data: AnalyticsWorkbookData, chart: AnalyticsChartData) -> None:
    if chart.visual_type == 'line':
        _add_line_chart(sheet, data, chart)
        return

    end_row = _SUPPORTING_DATA_ROW + min(len(chart.rows), 8) - 1

    if chart.visual_type == 'comparison_bar':
        visual = BarChart()
        visual.type = 'bar'
        visual.grouping = 'clustered'
        visual.overlap = 0
        visual.varyColors = False
        visual.title = chart.title
        visual.height = 8.0
        cast(Any, visual).width = 13.5
        visual.gapWidth = 60
        values = Reference(
            sheet,
            min_col=_SUPPORTING_START_COLUMN + 1,
            max_col=_SUPPORTING_START_COLUMN + 2,
            min_row=_SUPPORTING_HEADER_DATA_ROW,
            max_row=end_row,
        )
        categories = Reference(
            sheet,
            min_col=_SUPPORTING_START_COLUMN,
            min_row=_SUPPORTING_DATA_ROW,
            max_row=end_row,
        )
        visual.add_data(values, titles_from_data=True, from_rows=False)
        visual.set_categories(categories)
        visual.legend.position = 'b'
        visual.legend.overlay = False
        visual.dLbls = DataLabelList(
            showLegendKey=False,
            showVal=True,
            showCatName=False,
            showSerName=False,
            showPercent=False,
            showBubbleSize=False,
            showLeaderLines=False,
            dLblPos='outEnd',
            numFmt='0.0"%"',
        )
        visual.x_axis.scaling.orientation = 'maxMin'
        visual.x_axis.tickLblPos = 'nextTo'
        visual.x_axis.delete = False
        visual.y_axis.axPos = 'b'
        visual.y_axis.tickLblPos = 'nextTo'
        visual.y_axis.scaling.min = 0
        visual.y_axis.scaling.max = 100
        visual.y_axis.majorUnit = 25
        visual.y_axis.majorGridlines = None
        visual.y_axis.numFmt = '0"%"'
        visual.y_axis.title = None
        labels = ('Current', 'Projected')
        fills = (KAORI_ACCENT, KAORI_PRIMARY)
        lines = (KAORI_PRIMARY_SOFT, KAORI_PRIMARY_DARK)
        for index, series in enumerate(visual.series[:2]):
            series.tx = SeriesLabel(v=labels[index])
            series.graphicalProperties.solidFill = fills[index]
            series.graphicalProperties.line.solidFill = lines[index]
        _remove_frame(visual)
        sheet.add_chart(visual, 'A15')
        return

    if chart.visual_type == 'capacity_bar':
        visual = BarChart()
        visual.type = 'bar'
        visual.grouping = 'stacked'
        visual.overlap = 100
        visual.varyColors = False
        visual.title = chart.title
        visual.height = 8.0
        cast(Any, visual).width = 13.5
        visual.gapWidth = 70
        values = Reference(
            sheet,
            min_col=_SUPPORTING_START_COLUMN + 1,
            max_col=_SUPPORTING_START_COLUMN + 2,
            min_row=_SUPPORTING_HEADER_DATA_ROW,
            max_row=end_row,
        )
        categories = Reference(
            sheet,
            min_col=_SUPPORTING_START_COLUMN,
            min_row=_SUPPORTING_DATA_ROW,
            max_row=end_row,
        )
        visual.add_data(values, titles_from_data=True, from_rows=False)
        visual.set_categories(categories)
        visual.legend.position = 'b'
        visual.legend.overlay = False
        visual.dLbls = DataLabelList(
            showLegendKey=False,
            showVal=False,
            showCatName=False,
            showSerName=False,
            showPercent=False,
            showBubbleSize=False,
            showLeaderLines=False,
        )
        visual.x_axis.scaling.orientation = 'maxMin'
        visual.x_axis.delete = False
        visual.x_axis.tickLblPos = 'nextTo'
        visual.x_axis.title = None
        visual.y_axis.scaling.min = 0
        visual.y_axis.scaling.max = 100
        visual.y_axis.majorUnit = 25
        visual.y_axis.majorGridlines = None
        visual.y_axis.numFmt = '0"%"'
        visual.y_axis.title = None
        labels = ('Utilization', 'Headroom')
        fills = (KAORI_PRIMARY, KAORI_PRIMARY_SOFT)
        lines = (KAORI_PRIMARY_DARK, KAORI_ACCENT)
        for index, series in enumerate(visual.series[:2]):
            series.tx = SeriesLabel(v=labels[index])
            series.graphicalProperties.solidFill = fills[index]
            series.graphicalProperties.line.solidFill = lines[index]
        _remove_frame(visual)
        sheet.add_chart(visual, 'A15')
        return

    values = Reference(sheet, min_col=12, min_row=_SUPPORTING_HEADER_DATA_ROW, max_row=end_row)
    categories = Reference(sheet, min_col=11, min_row=_SUPPORTING_DATA_ROW, max_row=end_row)

    if chart.visual_type == 'ranked_bar':
        visual = BarChart()
        visual.type = 'bar'
        visual.grouping = 'clustered'
        visual.varyColors = False
        visual.visible_cells_only = False
        visual.title = chart.title
        visual.height = 8.0
        cast(Any, visual).width = 13.5
        visual.gapWidth = 65
        visual.add_data(values, titles_from_data=True)
        chart_categories = Reference(
            sheet,
            min_col=_CHART_CATEGORY_COLUMN,
            min_row=_SUPPORTING_DATA_ROW,
            max_row=end_row,
        )
        visual.set_categories(chart_categories)
        visual.legend = None
        data_labels = DataLabelList(
            showLegendKey=False,
            showVal=True,
            showCatName=False,
            showSerName=False,
            showPercent=False,
            showBubbleSize=False,
            showLeaderLines=False,
            dLblPos='outEnd',
        )
        visual.dLbls = data_labels
        if chart.series:
            data_labels.numFmt = _series_number_format(chart.series[0].number_format, chart.series[0].unit)
        visual.x_axis.scaling.orientation = 'maxMin'
        visual.x_axis.tickLblPos = 'nextTo'
        visual.x_axis.delete = False
        visual.y_axis.scaling.min = 0
        max_value = max((_number(row.get('value')) for row in chart.rows[:8]), default=0.0)
        if max_value > 0:
            visual.y_axis.scaling.max = max(max_value * 1.20, max_value + 1.0)
        visual.y_axis.title = None
        visual.y_axis.majorGridlines = None
        visual.series[0].graphicalProperties.solidFill = KAORI_PRIMARY
        visual.series[0].graphicalProperties.line.solidFill = KAORI_PRIMARY_DARK
        _remove_frame(visual)
        sheet.add_chart(visual, 'A15')
        return

    if chart.visual_type == 'composition_bar':
        share_values = Reference(sheet, min_col=13, min_row=_SUPPORTING_HEADER_DATA_ROW, max_row=end_row)
        visual = BarChart()
        visual.type = 'bar'
        visual.grouping = 'clustered'
        visual.varyColors = False
        visual.title = chart.title
        visual.height = 8.0
        cast(Any, visual).width = 13.5
        visual.gapWidth = 70
        visual.add_data(share_values, titles_from_data=True)
        visual.set_categories(categories)
        visual.legend = None
        visual.dLbls = DataLabelList(
            showLegendKey=False,
            showVal=True,
            showCatName=False,
            showSerName=False,
            showPercent=False,
            showBubbleSize=False,
            showLeaderLines=False,
            dLblPos='outEnd',
            numFmt='0.0"%"',
        )
        visual.x_axis.scaling.orientation = 'maxMin'
        visual.x_axis.tickLblPos = 'nextTo'
        visual.x_axis.delete = False
        visual.y_axis.scaling.min = 0
        visual.y_axis.scaling.max = 100
        visual.y_axis.majorUnit = 25
        visual.y_axis.numFmt = '0"%"'
        visual.y_axis.title = None
        visual.y_axis.majorGridlines = None
        points: list[DataPoint] = []
        for index, row in enumerate(chart.rows[:8]):
            point = DataPoint(idx=index)
            point.graphicalProperties.solidFill = _semantic_composition_color(
                str(row.get(chart.category_key) or ''), index
            )
            points.append(point)
        visual.series[0].data_points = points
        _remove_frame(visual)
        sheet.add_chart(visual, 'A15')
        return

    raise ValueError(f'Unsupported Analytics visual type: {chart.visual_type}.')


def _add_line_chart(sheet: Worksheet, data: AnalyticsWorkbookData, chart: AnalyticsChartData) -> None:
    source = data.table(chart.source_sheet)
    if source is None or not source.rows:
        return
    workbook = sheet.parent
    if workbook is None:
        raise ValueError('Analytics Overview must belong to a workbook.')
    source_sheet = cast(Worksheet, workbook[chart.source_sheet])
    column_indexes = {column.key: index for index, column in enumerate(source.columns, start=1)}
    category_index = column_indexes.get(chart.category_key)
    if category_index is None:
        return
    max_row = 5 + len(source.rows)
    visual = LineChart()
    visual.title = chart.title
    visual.height = 8.0
    cast(Any, visual).width = 13.5
    legend = visual.legend
    if legend is not None:
        legend.position = 'b'
        legend.overlay = False
    visual.x_axis.title = None
    visual.x_axis.tickLblPos = 'low'
    cast(Any, visual.x_axis).tickLblSkip = max(1, len(source.rows) // 6)
    percent_series = bool(chart.series) and all(series.unit == '%' for series in chart.series)
    if percent_series:
        visual.y_axis.title = None
        visual.y_axis.scaling.min = 0
        visual.y_axis.scaling.max = 100
        visual.y_axis.majorUnit = 25
        visual.y_axis.numFmt = '0"%"'
        visual.y_axis.majorGridlines = None
    else:
        visual.y_axis.title = chart.value_axis_title or 'Value'
        visual.y_axis.scaling.min = 0
    categories = Reference(source_sheet, min_col=category_index, min_row=6, max_row=max_row)
    visual.set_categories(categories)
    for index, series in enumerate(chart.series):
        column_index = column_indexes.get(series.key)
        if column_index is None:
            continue
        values = Reference(source_sheet, min_col=column_index, min_row=5, max_row=max_row)
        visual.add_data(values, titles_from_data=True, from_rows=False)
        rendered_series = visual.series[-1]
        rendered_series.tx = SeriesLabel(v=series.label)
        rendered_series.graphicalProperties.line.solidFill = _LINE_COLORS[index % len(_LINE_COLORS)]
        rendered_series.graphicalProperties.line.width = 24000
        if 'forecast' in series.label.casefold():
            rendered_series.graphicalProperties.line.prstDash = 'dash'
        rendered_series.marker.symbol = 'none'
    _remove_frame(visual)
    sheet.add_chart(visual, 'A15')


def _write_navigation(sheet: Worksheet) -> None:
    sheet.merge_cells(start_row=_NAV_HEADER_ROW, start_column=1, end_row=_NAV_HEADER_ROW, end_column=14)
    header = sheet.cell(row=_NAV_HEADER_ROW, column=1, value='Workbook Navigation')
    style_section_header(header)
    sheet.row_dimensions[_NAV_HEADER_ROW].height = 22

    workbook = sheet.parent
    if workbook is None:
        raise ValueError('Analytics Overview must belong to a workbook.')
    targets = [name for name in workbook.sheetnames if name != 'Overview'][:7]
    rows = (targets[:4], targets[4:7])
    spans = (
        ((1, 3), (4, 6), (7, 9), (10, 14)),
        ((1, 4), (5, 8), (9, 14)),
    )
    for row_offset, row_targets in enumerate(rows):
        row_index = _NAV_ROW + row_offset
        for target, (start_col, end_col) in zip(row_targets, spans[row_offset]):
            sheet.merge_cells(start_row=row_index, start_column=start_col, end_row=row_index, end_column=end_col)
            cell = sheet.cell(row=row_index, column=start_col, value=f'Open {target} →')
            cell.hyperlink = f"#'{target}'!A1"
            style_navigation_link(cell)
        sheet.row_dimensions[row_index].height = 26


def _merged_card(sheet: Worksheet, row: int, start: int, end: int, value: Any, style: Any) -> None:
    sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    cell = sheet.cell(row=row, column=start, value=value)
    style(cell)


def _overview_title(data: AnalyticsWorkbookData) -> str:
    report_type = data.profile.report_type
    if report_type == 'analytics_overview':
        return 'KAORI — Analytics Overview'
    if report_type == 'object_storage_analysis':
        return 'KAORI — Object Storage Analytics'
    if data.profile.title.endswith(' Analysis'):
        return f'KAORI — {data.profile.title.removesuffix(" Analysis")} Analytics'
    return f'KAORI — {data.profile.title} Analytics'


def _context_line(data: AnalyticsWorkbookData) -> str:
    cutoff = _display(data.context.data_cutoff)
    return (
        f'Reference: {data.context.report_id}  |  {data.context.scope_label}: {data.context.scope_value}  |  '
        f'Namespace: {data.context.namespace_scope}  |  Period: {data.context.period_label}  |  Data Cut-off: {cutoff}'
    )


def _series_number_format(number_format: str, unit: str | None) -> str:
    rendered_unit = str(unit or '').strip()
    if rendered_unit.casefold() in {'', 'n/a', 'na', 'not available', 'unknown', '-'}:
        return number_format
    safe_unit = rendered_unit.replace('"', '')
    return f'{number_format}" {safe_unit}"'


def _short_chart_label(value: str, *, max_length: int = 28) -> str:
    rendered = ' '.join(str(value or '').split())
    if len(rendered) <= max_length:
        return rendered
    return f'{rendered[: max_length - 1].rstrip()}…'


def _semantic_composition_color(category: str, index: int) -> str:
    key = category.casefold()
    if 'critical' in key or 'not covered' in key or 'failed' in key:
        return KAORI_RED_TEXT
    if 'warning' in key or 'pending' in key:
        return KAORI_AMBER_TEXT
    if 'info' in key:
        return KAORI_BLUE_TEXT
    if 'covered' in key or 'healthy' in key or 'normal' in key:
        return KAORI_GREEN_TEXT
    return _DOUGHNUT_COLORS[index % len(_DOUGHNUT_COLORS)]


def _remove_frame(chart: BarChart | LineChart) -> None:
    chart.graphical_properties = GraphicalProperties()
    chart.graphical_properties.noFill = True
    chart.graphical_properties.line.noFill = True
    chart.graphical_properties.line.prstDash = None


def _number(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _display(value: Any) -> str:
    if value in (None, ''):
        return 'Not Available'
    return str(value)
