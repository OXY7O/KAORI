from __future__ import annotations

from typing import Any, Final, Literal

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins, PrintPageSetup
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

KAORI_PRIMARY: Final = '7C3AED'
KAORI_PRIMARY_DARK: Final = '5B21B6'
KAORI_ACCENT: Final = 'A78BFA'
KAORI_PRIMARY_SOFT: Final = 'EDE9FE'
KAORI_PRIMARY_FAINT: Final = 'F5F3FF'
KAORI_INK: Final = '111827'
KAORI_MUTED: Final = '64748B'
KAORI_LINE: Final = 'CBD5E1'
KAORI_LINE_LIGHT: Final = 'E2E8F0'
KAORI_PANEL: Final = 'F8FAFC'
KAORI_GREEN_FILL: Final = 'DCFCE7'
KAORI_GREEN_TEXT: Final = '166534'
KAORI_AMBER_FILL: Final = 'FEF3C7'
KAORI_AMBER_TEXT: Final = '92400E'
KAORI_RED_FILL: Final = 'FEE2E2'
KAORI_RED_TEXT: Final = '991B1B'
KAORI_BLUE_FILL: Final = 'DBEAFE'
KAORI_BLUE_TEXT: Final = '1E40AF'
KAORI_SLATE_FILL: Final = 'F1F5F9'
KAORI_SLATE_TEXT: Final = '475569'
KAORI_WHITE: Final = 'FFFFFF'

BODY_FONT: Final = 'Aptos'
TITLE_FONT: Final = 'Aptos Display'

THIN_LINE: Final = Side(style='thin', color=KAORI_LINE_LIGHT)
MEDIUM_PRIMARY_LINE: Final = Side(style='medium', color=KAORI_PRIMARY)
HEADER_BORDER: Final = Border(bottom=Side(style='thin', color=KAORI_PRIMARY_DARK))
BODY_BORDER: Final = Border(bottom=THIN_LINE)
CARD_BORDER: Final = Border(
    left=Side(style='thin', color=KAORI_LINE_LIGHT),
    right=Side(style='thin', color=KAORI_LINE_LIGHT),
    top=Side(style='thin', color=KAORI_LINE_LIGHT),
    bottom=Side(style='thin', color=KAORI_LINE_LIGHT),
)


def style_title(cell: Any) -> None:
    cell.font = Font(name=TITLE_FONT, size=17, bold=True, color=KAORI_WHITE)
    cell.fill = PatternFill('solid', fgColor=KAORI_PRIMARY_DARK)
    cell.alignment = Alignment(horizontal='left', vertical='center')


def style_subtitle(cell: Any) -> None:
    cell.font = Font(name=BODY_FONT, size=10, color=KAORI_INK)
    cell.fill = PatternFill('solid', fgColor=KAORI_PRIMARY_SOFT)
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def style_context(cell: Any) -> None:
    cell.font = Font(name=BODY_FONT, size=9, color=KAORI_MUTED)
    cell.fill = PatternFill('solid', fgColor=KAORI_PANEL)
    cell.border = Border(bottom=THIN_LINE)
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def style_header(cell: Any) -> None:
    cell.font = Font(name=BODY_FONT, size=9, bold=True, color=KAORI_WHITE)
    cell.fill = PatternFill('solid', fgColor=KAORI_PRIMARY)
    cell.border = HEADER_BORDER
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def style_section_header(cell: Any) -> None:
    cell.font = Font(name=BODY_FONT, size=10, bold=True, color=KAORI_PRIMARY_DARK)
    cell.fill = PatternFill('solid', fgColor=KAORI_PRIMARY_SOFT)
    cell.border = Border(bottom=MEDIUM_PRIMARY_LINE)
    cell.alignment = Alignment(horizontal='left', vertical='center')


def style_body(
    cell: Any,
    *,
    alternate: bool = False,
    horizontal: Literal['left', 'center', 'right'] = 'left',
    wrap_text: bool = True,
) -> None:
    cell.font = Font(name=BODY_FONT, size=9, color=KAORI_INK)
    cell.border = BODY_BORDER
    cell.alignment = Alignment(horizontal=horizontal, vertical='center', wrap_text=wrap_text)
    if alternate:
        cell.fill = PatternFill('solid', fgColor=KAORI_PANEL)


def style_empty(cell: Any, *, primary: bool) -> None:
    cell.font = Font(name=BODY_FONT, size=9, italic=primary, color=KAORI_MUTED)
    cell.fill = PatternFill('solid', fgColor=KAORI_PANEL)
    cell.border = BODY_BORDER
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def style_summary_label(cell: Any) -> None:
    cell.font = Font(name=BODY_FONT, size=9, bold=True, color=KAORI_MUTED)
    cell.fill = PatternFill('solid', fgColor=KAORI_PANEL)
    cell.border = CARD_BORDER
    cell.alignment = Alignment(horizontal='left', vertical='center')


def style_summary_value(cell: Any) -> None:
    cell.font = Font(name=TITLE_FONT, size=15, bold=True, color=KAORI_INK)
    cell.fill = PatternFill('solid', fgColor=KAORI_WHITE)
    cell.border = CARD_BORDER
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def style_readout(cell: Any) -> None:
    cell.font = Font(name=BODY_FONT, size=10, color=KAORI_INK)
    cell.fill = PatternFill('solid', fgColor=KAORI_PRIMARY_FAINT)
    cell.border = Border(left=MEDIUM_PRIMARY_LINE, bottom=THIN_LINE)
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)


def style_navigation_link(cell: Any) -> None:
    cell.font = Font(name=BODY_FONT, size=9, bold=True, color=KAORI_PRIMARY, underline='single')
    cell.fill = PatternFill('solid', fgColor=KAORI_PRIMARY_FAINT)
    cell.border = CARD_BORDER
    cell.alignment = Alignment(horizontal='left', vertical='center')


def style_metadata_section(cell: Any) -> None:
    cell.font = Font(name=BODY_FONT, size=9, bold=True, color=KAORI_PRIMARY_DARK)
    cell.fill = PatternFill('solid', fgColor=KAORI_PRIMARY_SOFT)
    cell.border = Border(bottom=Side(style='thin', color=KAORI_ACCENT))
    cell.alignment = Alignment(horizontal='left', vertical='center')


def configure_sheet_page(
    sheet: Worksheet,
    *,
    print_area: str,
    print_title_rows: str,
    orientation: Literal['landscape', 'portrait'] = 'landscape',
    fit_to_width: int = 1,
) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.print_area = print_area
    sheet.print_title_rows = print_title_rows
    sheet.page_setup = PrintPageSetup(
        worksheet=sheet,
        orientation=orientation,
        paperSize=sheet.PAPERSIZE_A4,
        fitToWidth=max(1, fit_to_width),
        fitToHeight=0,
    )
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_margins = PageMargins(
        left=0.25,
        right=0.25,
        top=0.45,
        bottom=0.45,
        header=0.2,
        footer=0.2,
    )


def set_footer(sheet: Worksheet, document_reference: str) -> None:
    footer = sheet.oddFooter
    if footer is None:
        return
    left = footer.left
    right = footer.right
    if left is not None:
        left.text = document_reference
        left.size = 8
        left.color = KAORI_MUTED
    if right is not None:
        right.text = 'Page &P of &N'
        right.size = 8
        right.color = KAORI_MUTED
