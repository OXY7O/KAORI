from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.export.xlsx.analytics_schema import (
    ANALYTICS_DICTIONARY_SHEET,
    ANALYTICS_METADATA_SHEET,
    ANALYTICS_OVERVIEW_SHEET,
    ANALYTICS_WORKBOOK_FAMILY,
    ANALYTICS_WORKBOOK_STANDARD,
)
from app.reports.export.xlsx.schema import XLSX_REPORTING_STANDARD


def validate_analytics_xlsx_artifact(output_path: Path) -> None:
    workbook = load_workbook(output_path, read_only=False, data_only=False)
    _validate_sheet_order(workbook.sheetnames)
    metadata_values = _metadata_values(cast(Worksheet, workbook[ANALYTICS_METADATA_SHEET]))
    _validate_metadata(metadata_values)

    overview = cast(Worksheet, workbook[ANALYTICS_OVERVIEW_SHEET])
    _validate_overview_structure(overview)
    _validate_navigation(workbook, overview)
    _validate_primary_visual(overview, str(metadata_values.get('Primary Visual') or '').strip())


def _validate_sheet_order(sheetnames: list[str]) -> None:
    if not sheetnames or sheetnames[0] != ANALYTICS_OVERVIEW_SHEET:
        raise ValueError('Analytics XLSX artifact must begin with Overview.')
    if sheetnames[-2:] != [ANALYTICS_METADATA_SHEET, ANALYTICS_DICTIONARY_SHEET]:
        raise ValueError('Analytics XLSX artifact must end with Metadata and Field Dictionary.')


def _metadata_values(metadata: Worksheet) -> dict[Any, Any]:
    return {
        metadata.cell(row=row, column=1).value: metadata.cell(row=row, column=2).value
        for row in range(1, metadata.max_row + 1)
        if metadata.cell(row=row, column=1).value
    }


def _validate_metadata(metadata_values: Mapping[Any, Any]) -> None:
    if metadata_values.get('Workbook Standard') != ANALYTICS_WORKBOOK_STANDARD:
        raise ValueError('Analytics XLSX artifact uses a stale or missing workbook standard.')
    if metadata_values.get('XLSX Standard') != XLSX_REPORTING_STANDARD:
        raise ValueError('Analytics XLSX artifact uses a stale or missing global XLSX standard.')
    if metadata_values.get('Report Family') != ANALYTICS_WORKBOOK_FAMILY:
        raise ValueError('Analytics XLSX artifact metadata must declare the analytics family.')


def _validate_overview_structure(overview: Worksheet) -> None:
    if not str(overview['A1'].value or '').startswith('KAORI — '):
        raise ValueError('Analytics XLSX Overview title is missing.')
    if overview['A9'].value != 'Executive Readout':
        raise ValueError('Analytics XLSX Overview must contain the Executive Readout section.')
    if overview['A13'].value != 'Primary Visual':
        raise ValueError('Analytics XLSX Overview must reserve the Primary Visual surface.')
    if overview['J13'].value != 'Supporting Data':
        raise ValueError('Analytics XLSX Overview must reserve the Supporting Data surface.')


def _validate_navigation(workbook: Any, overview: Worksheet) -> None:
    navigation_targets = {
        str(cell.hyperlink.target)
        for row in overview.iter_rows(min_row=32, max_row=33)
        for cell in row
        if cell.hyperlink is not None
    }
    for sheet_name in workbook.sheetnames[1:]:
        if f"#'{sheet_name}'!A1" not in navigation_targets:
            raise ValueError(f'Analytics XLSX Overview is missing navigation to {sheet_name!r}.')

    for sheet_name in workbook.sheetnames[1:-2]:
        sheet = cast(Worksheet, workbook[sheet_name])
        if sheet['A3'].hyperlink is None or sheet['A3'].hyperlink.target != "#'Overview'!A1":
            raise ValueError(f'Analytics XLSX sheet {sheet_name!r} must link back to Overview.')


def _validate_primary_visual(overview: Worksheet, primary_visual: str) -> None:
    overview_charts = cast(Any, overview)._charts
    if primary_visual and primary_visual != 'No chartable evidence' and not overview_charts:
        raise ValueError('Analytics XLSX Primary Visual is missing for chartable committed evidence.')
    if not overview_charts:
        return

    chart = overview_charts[0]
    chart_name = chart.__class__.__name__
    if chart_name == 'DoughnutChart':
        raise ValueError('Analytics XLSX no longer permits doughnut charts; use the global composition_bar contract.')
    if chart_name == 'BarChart':
        _validate_bar_chart(chart, overview, primary_visual)
    if chart_name == 'LineChart':
        _validate_line_chart(chart)


def _validate_bar_chart(chart: Any, overview: Worksheet, primary_visual: str) -> None:
    if getattr(chart, 'varyColors', None) is not False:
        raise ValueError('Analytics ranked/capacity bar charts must use controlled series colors.')

    series_count = len(chart.series)
    if series_count == 1 and chart.legend is not None:
        raise ValueError('Analytics single-series ranked bar charts must not use a redundant legend.')

    if primary_visual == 'ranked_bar':
        _validate_ranked_bar(chart, overview)
    if primary_visual == 'composition_bar':
        _validate_composition_bar(chart, series_count)
    if primary_visual == 'capacity_bar':
        _validate_capacity_bar(chart, series_count)
    if primary_visual == 'comparison_bar':
        _validate_comparison_bar(chart, series_count)


def _validate_ranked_bar(chart: Any, overview: Worksheet) -> None:
    data_labels = getattr(chart, 'dLbls', None)
    if (
        data_labels is None
        or data_labels.showVal is not True
        or data_labels.showLegendKey is not False
        or data_labels.showCatName is not False
        or data_labels.showSerName is not False
        or data_labels.showPercent is not False
        or data_labels.showBubbleSize is not False
        or data_labels.dLblPos != 'outEnd'
    ):
        raise ValueError('Analytics ranked bars require value-only direct labels for Excel compatibility.')
    if (
        chart.x_axis.delete is not False
        or chart.x_axis.tickLblPos != 'nextTo'
        or chart.y_axis.scaling.min != 0
        or chart.y_axis.title is not None
        or chart.y_axis.majorGridlines is not None
        or chart.visible_cells_only is not False
    ):
        raise ValueError('Analytics ranked bars require visible categories and a sparse zero-baseline value axis.')
    if not overview.column_dimensions['O'].hidden:
        raise ValueError('Analytics ranked bars require the hidden short-category helper column.')
    if (overview.column_dimensions['K'].width or 0) < 38 or (overview.column_dimensions['L'].width or 0) < 20:
        raise ValueError('Analytics ranked Supporting Data requires professional category/value column widths.')

    ranked_values: list[float] = []
    for row in range(16, min(overview.max_row, 23) + 1):
        short_label = str(overview.cell(row=row, column=15).value or '')
        if len(short_label) > 28:
            raise ValueError('Analytics ranked chart helper labels must be shortened without changing Supporting Data.')
        if 'not available' in str(overview.cell(row=row, column=12).number_format or '').casefold():
            raise ValueError('Analytics ranked numeric formats must not publish placeholder text as a unit.')
        value = overview.cell(row=row, column=12).value
        if isinstance(value, (int, float)):
            ranked_values.append(float(value))

    if ranked_values:
        required_max = max(ranked_values) * 1.10
        axis_max = chart.y_axis.scaling.max
        if axis_max is None or axis_max < required_max:
            raise ValueError('Analytics ranked bars must reserve value-axis headroom for direct labels.')


def _validate_composition_bar(chart: Any, series_count: int) -> None:
    if series_count != 1 or chart.grouping != 'clustered' or chart.legend is not None:
        raise ValueError('Analytics composition bars require one legend-free clustered share series.')
    data_labels = getattr(chart, 'dLbls', None)
    if (
        data_labels is None
        or data_labels.showVal is not True
        or data_labels.showLegendKey is not False
        or data_labels.showCatName is not False
        or data_labels.showSerName is not False
        or data_labels.showPercent is not False
        or data_labels.showBubbleSize is not False
        or data_labels.dLblPos != 'outEnd'
        or data_labels.numFmt != '0.0"%"'
    ):
        raise ValueError('Analytics composition bars require percentage-only value labels.')
    if (
        chart.x_axis.delete is not False
        or chart.x_axis.tickLblPos != 'nextTo'
        or chart.y_axis.scaling.min != 0
        or chart.y_axis.scaling.max != 100
        or chart.y_axis.majorUnit != 25
        or chart.y_axis.title is not None
        or chart.y_axis.majorGridlines is not None
    ):
        raise ValueError('Analytics composition bars require visible categories and a clean 0-100 percent scale.')


def _validate_capacity_bar(chart: Any, series_count: int) -> None:
    if series_count != 2 or chart.grouping != 'stacked' or chart.overlap != 100:
        raise ValueError('Analytics capacity bars require a 100% stacked utilization/headroom presentation.')
    if chart.legend is None or chart.legend.position != 'b':
        raise ValueError('Analytics capacity bars require a bottom legend.')
    labels = [str(getattr(series.tx, 'v', '') or '').strip() for series in chart.series]
    if labels != ['Utilization', 'Headroom']:
        raise ValueError('Analytics capacity bars require Utilization and Headroom series labels.')

    data_labels = getattr(chart, 'dLbls', None)
    if (
        data_labels is None
        or data_labels.showVal is not False
        or data_labels.showLegendKey is not False
        or data_labels.showCatName is not False
        or data_labels.showSerName is not False
        or data_labels.showPercent is not False
    ):
        raise ValueError(
            'Analytics capacity bars must keep direct data labels disabled to prevent Excel text collisions.'
        )
    if (
        chart.x_axis.delete is not False
        or chart.x_axis.tickLblPos != 'nextTo'
        or chart.x_axis.title is not None
        or chart.y_axis.scaling.min != 0
        or chart.y_axis.scaling.max != 100
        or chart.y_axis.majorUnit != 25
        or chart.y_axis.majorGridlines is not None
        or chart.y_axis.title is not None
    ):
        raise ValueError('Analytics capacity bars require visible resource categories and a clean 0-100 percent scale.')


def _validate_comparison_bar(chart: Any, series_count: int) -> None:
    if series_count != 2 or chart.grouping != 'clustered':
        raise ValueError('Analytics comparison bars require clustered Current/Projected series.')
    if chart.legend is None or chart.legend.position != 'b':
        raise ValueError('Analytics comparison bars require a bottom legend.')
    labels = [str(getattr(series.tx, 'v', '') or '').strip() for series in chart.series]
    if labels != ['Current', 'Projected']:
        raise ValueError('Analytics comparison bars require Current and Projected series labels.')

    data_labels = getattr(chart, 'dLbls', None)
    if (
        data_labels is None
        or data_labels.showVal is not True
        or data_labels.showLegendKey is not False
        or data_labels.showCatName is not False
        or data_labels.showSerName is not False
        or data_labels.showPercent is not False
        or data_labels.showBubbleSize is not False
        or data_labels.dLblPos != 'outEnd'
    ):
        raise ValueError('Analytics comparison bars require value-only data labels for Excel compatibility.')
    if (
        chart.x_axis.tickLblPos != 'nextTo'
        or chart.x_axis.delete is not False
        or chart.y_axis.title is not None
        or chart.y_axis.scaling.min != 0
        or chart.y_axis.scaling.max != 100
        or chart.y_axis.majorUnit != 25
        or chart.y_axis.majorGridlines is not None
    ):
        raise ValueError('Analytics comparison bars require the executive axis presentation contract.')


def _validate_line_chart(chart: Any) -> None:
    if chart.legend is None or chart.legend.position != 'b':
        raise ValueError('Analytics line charts require a bottom legend.')
    if any(not str(getattr(series.tx, 'v', '') or '').strip() for series in chart.series):
        raise ValueError('Analytics line charts require explicit series labels.')
    if (
        chart.x_axis.title is not None
        or chart.y_axis.title is not None
        or chart.y_axis.scaling.min != 0
        or chart.y_axis.scaling.max != 100
        or chart.y_axis.majorUnit != 25
        or chart.y_axis.majorGridlines is not None
        or any(getattr(series.marker, 'symbol', None) not in {None, 'none'} for series in chart.series)
    ):
        raise ValueError('Analytics line charts require the sparse percent presentation contract.')
